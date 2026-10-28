from pathlib import Path
import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn
from pagethaw.engine import inspect_pdf, convert_pdf
from pagethaw.fonts import read_font
from pagethaw.inline_forms import continuation_borders, attach_square
from pagethaw.editable import lines_on


def fixture(path, mismatch=False, marked=False, field=True):
    font = Path('C:/Windows/Fonts/arial.ttf')
    with pymupdf.open() as pdf:
        for n in range(2):
            p = pdf.new_page(width=500,height=500)
            p.insert_font(fontname='fixture',fontfile=font)
            p.insert_text((30,65), 'Yes        / No        agree: '+('__________' if field else 'ordinary text'), fontname='fixture',fontsize=10)
            chars = [c for b in p.get_text('rawdict')['blocks'] for l in b['lines'] for s in l['spans'] for c in s['chars']]
            x = chars[4]['bbox'][0]
            square = pymupdf.Rect(x,58,x+7,65)
            p.draw_rect(square)
            if marked:p.draw_line(square.tl,square.br)
            right = 470 if not mismatch or n == 0 else 460
            points = [(20,490),(20,40),(right,40),(right,490)] if n == 0 else [(20,10),(20,100),(right,100),(right,10)]
            p.draw_polyline(points,closePath=False)
        pdf.save(path)
    return read_font(font,installed=True)


@pytest.mark.parametrize('keep',[False,True])
@pytest.mark.parametrize('field',[False,True])
def test_inline_squares_and_paired_continuation_are_editable(tmp_path,keep,field):
    source=tmp_path/'forms.pdf';face=fixture(source,field=field)
    info=inspect_pdf(source);out=tmp_path/'forms.docx'
    result=convert_pdf(info,out,{n:face for n in info.fonts},mode='editable',keep_line_breaks=keep)
    assert result['content_check']=='passed'
    doc=Document(out)
    assert len(doc.tables)==1
    assert len(doc.tables[0].rows)>1
    assert len(doc.element.xpath('.//wp:inline'))==2
    assert not doc.element.xpath('.//wp:anchor')
    assert not list(doc.element.iter(qn('w:txbxContent')))
    assert not list(doc.element.iter(qn('w:trHeight')))
    text=''.join(t.text or '' for t in doc.element.iter(qn('w:t')))
    assert text.count('Yes')==2 and text.count('agree:')==2


def test_mismatched_continuation_not_paired(tmp_path):
    source=tmp_path/'bad.pdf';fixture(source,mismatch=True)
    with pymupdf.open(source) as pdf:assert not continuation_borders(pdf)


def test_marked_square_not_silently_emptied(tmp_path):
    source=tmp_path/'marked.pdf';fixture(source,marked=True)
    with pymupdf.open(source) as pdf:
        page=pdf[0];lines=lines_on(page)
        assert not attach_square(page,page.get_drawings()[0],lines)
