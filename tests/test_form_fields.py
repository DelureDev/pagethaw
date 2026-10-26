"""Regression checks for source field boundaries and date-year blanks."""
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn

from pagethaw.editable import blank_ranges, lines_on
from pagethaw.engine import inspect_pdf, convert_pdf
from pagethaw.fonts import read_font


@pytest.mark.parametrize('suffix', [' г.', 'г.', ' y.', 'y.', ''])
def test_only_date_year_underscores_join_day_month_fields(suffix):
    text = 'Date «______» __________ 20______'+suffix
    ranges = blank_ranges(text)
    assert len(ranges) == 3
    assert ''.join(text[a:b] for a, b in ranges) == '_'*22
    assert ''.join(c for i, c in enumerate(text) if not any(a <= i < b for a, b in ranges)) == 'Date «»  20'+suffix


@pytest.mark.parametrize('text', [
    'account_id a___b code____ 20____ 1234____',
    '____ ____ Item 20____', '____ 20____',
    '____ ____ 20____account', '____ ____ 20____x_other',
])
def test_identifiers_and_ambiguous_numeric_suffixes_stay_literal(text):
    assert all(text[a-1:a] != '0' for a, b in blank_ranges(text))


@pytest.mark.parametrize('row', ['______  ______  ', 'Text ______   ', '______    Text'])
def test_real_whitespace_is_not_drawn_as_part_of_a_blank(tmp_path, row):
    font = Path('C:/Windows/Fonts/arial.ttf')
    if not font.exists():
        pytest.skip('Windows Arial required for this fixture.')
    face = read_font(font, installed=True)
    src = tmp_path/'spaces.pdf'
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_font(fontname='fixture', fontfile=font)
        page.insert_text((30, 60), row, fontname='fixture', fontsize=10)
        pdf.save(src)
    info = inspect_pdf(src)
    output = tmp_path/'spaces.docx'
    convert_pdf(info, output, {n: face for n in info.fonts}, mode='editable')
    with pymupdf.open(src) as pdf:
        line = lines_on(pdf[0])[0]
    chars = [c for s in line.spans for c in s['chars']]
    expected = [(chars[a]['bbox'][0]-line.bbox.x0, chars[b-1]['bbox'][2]-line.bbox.x0)
                for a, b in blank_ranges(line.text)]
    table = Document(output).tables[0]
    actual, x = [], 0
    for col, cell in zip(table.columns, table.rows[0].cells):
        if cell._tc.xpath('./w:tcPr/w:tcBorders/w:bottom[@w:val="single"]'):
            actual.append((x, x+col.width.pt))
        x += col.width.pt
    assert len(actual) == len(expected)
    for a, e in zip(actual, expected):
        assert a == pytest.approx(e, abs=.2)


@pytest.mark.parametrize('keep_lines', [False, True])
def test_field_edges_and_date_content_survive_conversion(tmp_path, keep_lines):
    font = Path('C:/Windows/Fonts/arial.ttf')
    if not font.exists():
        pytest.skip('Windows Arial required for this fixture.')
    face = read_font(font, installed=True)
    src = tmp_path/'fields.pdf'
    row = 'Name: __________ Number: ______ Date «____» ______ 20____г.'
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_font(fontname='fixture', fontfile=font)
        page.insert_text((30, 60), row, fontname='fixture', fontsize=10)
        page.insert_text((30, 95), 'Ordinary text account_id a___b 20____.', fontname='fixture', fontsize=10)
        pdf.save(src)
    info = inspect_pdf(src)
    output = tmp_path/'fields.docx'
    result = convert_pdf(info, output, {n: face for n in info.fonts}, mode='editable', keep_line_breaks=keep_lines)
    assert result['content_check'] == 'passed'
    assert result['fields'] == 5
    doc = Document(output)
    table = doc.tables[0]
    with pymupdf.open(src) as pdf:
        line = lines_on(pdf[0])[0]
    chars = [c for s in line.spans for c in s['chars']]
    expected = [(chars[a]['bbox'][0]-line.bbox.x0, chars[b-1]['bbox'][2]-line.bbox.x0)
                for a, b in blank_ranges(line.text)]
    actual = []
    x = 0
    for col, cell in zip(table.columns, table.rows[0].cells):
        if cell._tc.xpath('./w:tcPr/w:tcBorders/w:bottom[@w:val="single"]'):
            actual.append((x, x+col.width.pt))
        x += col.width.pt
    assert len(actual) == len(expected)
    for a, e in zip(actual, expected):
        assert a == pytest.approx(e, abs=.2)
    text = ''.join(t.text or '' for t in doc.element.iter(qn('w:t')))
    assert '20' in text and 'г.' in text
    assert 'account_id a___b 20____.' in text
    assert not list(doc.element.iter(qn('w:txbxContent')))
    assert not list(doc.element.iter(qn('w:framePr')))
    assert not list(doc.element.iter(qn('w:trHeight')))
    field = table.rows[0].cells[1]
    field.paragraphs[0].runs[0].text = 'One\nTwo\nThree'
    doc.save(output)
    assert Document(output).tables[0].cell(0, 1).text == 'One\nTwo\nThree'
