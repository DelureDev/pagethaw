from pathlib import Path
import copy
import pymupdf
import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from pagethaw.editable import build_document, verify_document, UnsupportedLayout
from pagethaw.engine import inspect_pdf, convert_pdf
from pagethaw.fonts import read_font
from pagethaw.footnotes import read_part


@pytest.fixture
def face():
    path=Path('C:/Windows/Fonts/arial.ttf')
    if not path.exists():pytest.skip('Windows Arial required')
    return read_font(path,installed=True)


def fixture(tmp_path,face,separator='underscores',markers=('1',),anchors=None,raised_label=True):
    path=tmp_path/'notes.pdf'
    font=pymupdf.Font(fontfile=face.path)
    def row(page,y,label,prose,raised=True):
        page.insert_text((30,y-4 if raised else y),label,fontname='fixture',fontsize=7 if raised else 12)
        page.insert_text((30+font.text_length(label,fontsize=7 if raised else 12),y),' '+prose,fontname='fixture',fontsize=12)
    with pymupdf.open() as pdf:
        page=pdf.new_page(width=500,height=500)
        page.insert_font(fontname='fixture',fontfile=face.path)
        for i,marker in enumerate(markers if anchors is None else anchors):
            y=50+30*i
            prefix='Body text with reference'
            page.insert_text((30,y),prefix,fontname='fixture',fontsize=12)
            x=30+font.text_length(prefix,fontsize=12)
            page.insert_text((x,y-4),marker,fontname='fixture',fontsize=7)
            page.insert_text((x+font.text_length(marker,fontsize=7),y),'.',fontname='fixture',fontsize=12)
        page.insert_text((30,180),'Body remains before its continuation.',fontname='fixture',fontsize=12)
        if separator=='underscores':page.insert_text((30,390),'_'*40,fontname='fixture',fontsize=12)
        else:page.draw_line((30,390),(300,390))
        for i,marker in enumerate(markers):
            row(page,410+i*30,marker,'Arbitrary note content '+marker,raised_label)
            page.insert_text((30,422+i*30),'Wrapped note continuation.',fontname='fixture',fontsize=12)
        page=pdf.new_page(width=500,height=500)
        page.insert_font(fontname='fixture',fontfile=face.path)
        page.insert_text((30,40),'Body on following source page.',fontname='fixture',fontsize=12)
        pdf.save(path)
    info=inspect_pdf(path)
    assert not info.error
    return info,{name:face for name in info.fonts}


@pytest.mark.parametrize('separator',['underscores','vector'])
@pytest.mark.parametrize('keep_lines',[False,True])
@pytest.mark.parametrize('markers',[('7',),('1','2'),('†',)])
def test_native_footnotes_preserve_body_note_order_and_custom_marks(tmp_path,face,separator,keep_lines,markers):
    info,mapping=fixture(tmp_path,face,separator,markers)
    target=tmp_path/'notes.docx'
    convert_pdf(info,target,mapping,mode='editable',keep_line_breaks=keep_lines)
    doc=Document(target);root=read_part(doc)
    body=''.join(t.text or '' for t in doc.element.iter(qn('w:t')))
    assert 'Arbitrary note' not in body and 'Wrapped note' not in body
    assert body.index('before its continuation') < body.index('Body on following')
    continuation=next(p for p in doc.paragraphs if 'Body on following' in p.text)
    assert continuation._p.getprevious().xpath('.//w:br[@w:type="page"]')
    references=list(doc.element.iter(qn('w:footnoteReference')))
    assert len(references)==len(markers)
    notes=[n for n in root if n.get(qn('w:type')) is None]
    assert len(notes)==len(markers)
    for marker,ref,note in zip(markers,references,notes):
        assert ref.get(qn('w:id'))==note.get(qn('w:id'))
        assert ref.get(qn('w:customMarkFollows'))=='1'
        assert ref.getparent().find(qn('w:t')).text==marker
        text=''.join(t.text or '' for t in note.iter(qn('w:t')))
        assert text.startswith(marker) and 'Arbitrary note content '+marker in text
        assert 'Wrapped note continuation.' in text
        assert not list(note.iter(qn('w:txbxContent')))
        assert not list(note.iter(qn('w:framePr')))
    # Verifier must work after OPC deserialization, not only with an emitter ledger.
    with pymupdf.open(info.path) as pdf:_,expected,_=build_document(pdf,mapping,keep_line_breaks=keep_lines)
    verify_document(doc,expected)


@pytest.mark.parametrize('anchors',[(),('1','1')])
def test_missing_or_ambiguous_note_anchor_is_rejected(tmp_path,face,anchors):
    info,mapping=fixture(tmp_path,face,anchors=anchors)
    target=tmp_path/'invalid.docx'
    with pytest.raises(UnsupportedLayout,match='footnote references are ambiguous'):
        convert_pdf(info,target,mapping,mode='editable')
    assert not target.exists()


def test_ordinary_footer_and_body_exponent_are_not_promoted(tmp_path,face):
    info,mapping=fixture(tmp_path,face,raised_label=False)
    target=tmp_path/'footer.docx'
    convert_pdf(info,target,mapping,mode='editable')
    doc=Document(target)
    assert not list(doc.element.iter(qn('w:footnoteReference')))
    assert 'Arbitrary note content' in ''.join(p.text for p in doc.paragraphs)


@pytest.mark.parametrize('x,y',[(210,480),(30,480),(210,434)])
def test_separate_footer_is_not_silently_absorbed_into_note(tmp_path,face,x,y):
    info,mapping=fixture(tmp_path,face)
    with pymupdf.open(info.path) as pdf:
        pdf[0].insert_text((x,y),'Page 1 of 2',fontname='fixture',fontsize=10)
        pdf.saveIncr()
    info=inspect_pdf(info.path)
    with pytest.raises(UnsupportedLayout,match='footnote references are ambiguous'):
        convert_pdf(info,tmp_path/'footer.docx',mapping,mode='editable')
    assert not (tmp_path/'footer.docx').exists()


def test_two_note_columns_merged_into_one_pdf_line_are_rejected(tmp_path,face):
    path=tmp_path/'columns.pdf'
    font=pymupdf.Font(fontfile=face.path)
    with pymupdf.open() as pdf:
        page=pdf.new_page(width=500,height=500)
        page.insert_font(fontname='fixture',fontfile=face.path)
        for i,marker in enumerate(['1','2']):
            prefix='Body reference'
            page.insert_text((30,50+30*i),prefix,fontname='fixture',fontsize=12)
            page.insert_text((30+font.text_length(prefix,fontsize=12),46+30*i),marker,fontname='fixture',fontsize=7)
            x=30+250*i
            page.insert_text((x,406),marker,fontname='fixture',fontsize=7)
            page.insert_text((x+font.text_length(marker,fontsize=7),410),' Note '+marker,fontname='fixture',fontsize=12)
        page.draw_line((30,390),(470,390))
        pdf.save(path)
    info=inspect_pdf(path)
    with pytest.raises(UnsupportedLayout,match='footnote references are ambiguous'):
        convert_pdf(info,tmp_path/'columns.docx',{n:face for n in info.fonts},mode='editable')
    assert not (tmp_path/'columns.docx').exists()


@pytest.mark.parametrize('mutation',['delete','reorder','font','bold','reference','orphan','frame','move_anchor'])
def test_note_content_emphasis_identity_and_anchor_tampering_fails(tmp_path,face,mutation):
    info,mapping=fixture(tmp_path,face)
    with pymupdf.open(info.path) as pdf:doc,expected,_=build_document(pdf,mapping)
    part=doc.part.part_related_by(RT.FOOTNOTES)
    note=next(n for n in part.element if n.get(qn('w:type')) is None)
    runs=[r for r in note.iter(qn('w:r')) if r.find(qn('w:t')) is not None]
    if mutation=='delete':runs[-1].getparent().remove(runs[-1])
    elif mutation=='reorder':
        parent=runs[0].getparent();runs[-1].getparent().remove(runs[-1]);parent.insert(1,runs[-1])
    elif mutation=='font':runs[-1].find(qn('w:rPr')).find(qn('w:rFonts')).set(qn('w:ascii'),'Wrong font')
    elif mutation=='bold':runs[-1].find(qn('w:rPr')).find(qn('w:b')).set(qn('w:val'),'1')
    elif mutation=='reference':next(doc.element.iter(qn('w:footnoteReference'))).set(qn('w:id'),'99')
    elif mutation=='orphan':
        extra=copy.deepcopy(note);extra.set(qn('w:id'),'99');part.element.append(extra)
    elif mutation=='frame':note.append(OxmlElement('w:framePr'))
    else:
        ref=next(doc.element.iter(qn('w:footnoteReference')))
        run=ref.getparent();run.getparent().remove(run)
        doc.paragraphs[-1]._p.append(run)
    with pytest.raises(ValueError):verify_document(doc,expected)
