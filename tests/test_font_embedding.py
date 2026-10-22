from dataclasses import replace
import io
from pathlib import Path
import threading
import uuid
import zipfile

from docx import Document
from docx.oxml.ns import qn
from fontTools.fontBuilder import FontBuilder
from fontTools.pens.ttGlyphPen import TTGlyphPen
from fontTools.ttLib import TTFont
import pytest

from pagethaw.font_embedding import eligibility, embed_fonts, obfuscate, verify_embedded_fonts
from pagethaw.fonts import read_font


def make_font(path, flags=8, style='Regular', family='PageThaw Synthetic Test'):
    """Original geometric test glyphs; no system or proprietary font bytes."""
    builder = FontBuilder(1000, isTTF=True)
    cmap = {i: f'g{i}' for i in list(range(32, 127))+[1071]}
    names = ['.notdef', *cmap.values()]
    builder.setupGlyphOrder(names)
    builder.setupCharacterMap(cmap)
    glyphs = {}
    for i, name in enumerate(names):
        pen = TTGlyphPen(None)
        if name != 'g32':
            pen.moveTo((100, 0)); pen.lineTo((100, 700)); pen.lineTo((800, 700)); pen.lineTo((800, 500))
            pen.lineTo((350, 500)); pen.lineTo((350, 0)); pen.closePath()
        glyphs[name] = pen.glyph()
    builder.setupGlyf(glyphs)
    builder.setupHorizontalMetrics({n:(1000, 100) for n in names})
    builder.setupHorizontalHeader(ascent=800, descent=-200)
    builder.setupNameTable({'familyName': family, 'styleName': style, 'uniqueFontIdentifier': family+' '+style,
                           'fullName': family+' '+style, 'psName': family.replace(' ', '')+'-'+style})
    builder.setupOS2(sTypoAscender=800, sTypoDescender=-200, usWinAscent=800, usWinDescent=200, fsType=flags,
                     fsSelection=64, ulCodePageRange1=5, sxHeight=500, sCapHeight=700)
    builder.setupPost()
    builder.save(path)
    return read_font(path)


@pytest.mark.parametrize('flags,allowed', [(0,True), (8,True), (0x100,True), (0x108,True),
    (2,False), (4,False), (0x104,False), (0x208,False), (6,False), (12,False), (1,False), (0x400,False)])
def test_permission_flags_and_complete_bytes(tmp_path, flags, allowed):
    face = make_font(tmp_path/'test.ttf', flags)
    doc = Document()
    doc.add_paragraph('Test').runs[0].font.name = face.family
    result = embed_fonts(doc, [face, face])
    assert bool(result['embedded_fonts']) == allowed
    assert bool(result['font_embedding_skipped']) != allowed
    assert len(result['embedded_fonts'] or result['font_embedding_skipped']) == 1
    path = tmp_path/'test.docx'
    doc.save(path)
    verify_embedded_fonts(path, result)
    with zipfile.ZipFile(path) as z:
        parts = [n for n in z.namelist() if n.endswith('.odttf')]
        assert len(parts) == int(allowed)
        if allowed:
            from lxml import etree
            root = etree.fromstring(z.read('word/fontTable.xml'))
            entry = next(root.iter(qn('w:embedRegular')))
            data = obfuscate(z.read(parts[0]), entry.get(qn('w:fontKey')).strip('{}'))
            assert data == Path(face.path).read_bytes()  # full glyph set, unchanged permissions


def test_obfuscation_known_vector():
    key = uuid.UUID('00112233-4455-6677-8899-aabbccddeeff')
    assert obfuscate(bytes(32)+b'unchanged', key) == bytes.fromhex('ffeeddccbbaa99887766554433221100')*2+b'unchanged'


def test_missing_malformed_and_unsupported_fonts(tmp_path):
    face = make_font(tmp_path/'test.ttf')
    with TTFont(face.path) as font:
        del font['OS/2']
        buffer = io.BytesIO(); font.save(buffer)
    assert 'permission' in eligibility(buffer.getvalue())
    assert 'could not' in eligibility(b'broken')
    with TTFont(face.path) as font:
        font.flavor = 'woff'
        buffer = io.BytesIO(); font.save(buffer)
    assert 'TrueType' in eligibility(buffer.getvalue())


def test_four_styles_conflicts_and_changed_font(tmp_path):
    faces = [make_font(tmp_path/f'{style}.ttf', style=style) for style in ['Regular','Bold','Italic','BoldItalic']]
    # Font face metadata selects the exact Word style slot.
    faces = [replace(f, bold=bool(i&1), italic=bool(i&2)) for i,f in enumerate(faces)]
    doc = Document(); result = embed_fonts(doc, faces)
    assert {r['style'] for r in result['embedded_fonts']} == {'Regular','Bold','Italic','BoldItalic'}
    path=tmp_path/'test.docx'; doc.save(path); verify_embedded_fonts(path,result)
    conflict = replace(faces[1], bold=False)
    result = embed_fonts(Document(), [faces[0], conflict])
    assert not result['embedded_fonts'] and len(result['font_embedding_skipped']) == 2
    Path(faces[0].path).write_bytes(b'changed')
    with pytest.raises(ValueError, match='changed'):
        embed_fonts(Document(), [faces[0]])


def test_cancel_and_payload_tampering(tmp_path):
    face = make_font(tmp_path/'test.ttf')
    cancel = threading.Event(); cancel.set()
    with pytest.raises(InterruptedError):
        embed_fonts(Document(), [face], cancel)
    doc=Document(); result=embed_fonts(doc,[face]); path=tmp_path/'test.docx'; doc.save(path)
    result['embedded_fonts'][0]['sha256']='bad'
    with pytest.raises(ValueError,match='verification'):
        verify_embedded_fonts(path,result)
