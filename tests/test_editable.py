from pathlib import Path
import threading

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn

from pagethaw.editable import build_document, verify_document, UnsupportedLayout
from pagethaw.engine import inspect_pdf, convert_batch, convert_pdf
from pagethaw.fonts import read_font


@pytest.fixture
def face():
    path = Path('C:/Windows/Fonts/arial.ttf')
    if not path.exists():
        pytest.skip('Windows Arial is required for these fixtures.')
    return read_font(path, installed=True)


def source(tmp_path, face, draw=None):
    path = tmp_path/'form.pdf'
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_font(fontname='fixture', fontfile=face.path)
        if draw:
            draw(page)
        else:
            page.insert_text((30, 60), 'Name: __________________', fontname='fixture', fontsize=12)
            page.insert_text((30, 100), 'A paragraph with ordinary editable text.', fontname='fixture', fontsize=12)
        pdf.save(path)
    info = inspect_pdf(path)
    assert not info.error
    return info, {name: face for name in info.fonts}


def test_flowing_text_and_expanding_fields(tmp_path, face):
    info, mapping = source(tmp_path, face)
    target = tmp_path/'result.docx'
    result = convert_pdf(info, target, mapping, mode='editable')
    doc = Document(target)
    assert result['content_check'] == 'passed' and result['fields'] == 1
    assert result['source_pages'] == 1 and 'pages' not in result
    assert not list(doc.element.iter(qn('w:txbxContent')))
    assert not list(doc.element.iter(qn('w:framePr')))
    assert not list(doc.element.iter(qn('w:trHeight')))
    assert any('ordinary editable text' in p.text for p in doc.paragraphs)
    assert len(doc.tables) == 1
    field = doc.tables[0].cell(0, 1)
    field.paragraphs[0].runs[0].text = 'First line\nSecond line\nThird line'
    doc.save(tmp_path/'edited.docx')
    edited = Document(tmp_path/'edited.docx')
    assert 'Third line' in edited.tables[0].cell(0, 1).text
    assert any('ordinary editable text' in p.text for p in edited.paragraphs)


@pytest.mark.parametrize('keep_lines', [False, True])
def test_whitespace_only_source_row_does_not_break_margin_inference(tmp_path, face, keep_lines):
    def draw(page):
        for y, text in [(60, 'X'*70), (72, '   '), (84, 'continuation')]:
            page.insert_text((30, y), text, fontname='fixture', fontsize=10)
    info, mapping = source(tmp_path, face, draw)
    target = tmp_path/'blank-row.docx'
    convert_pdf(info, target, mapping, mode='editable', keep_line_breaks=keep_lines)
    doc = Document(target)
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert 'X'*70 in text and 'continuation' in text
    if keep_lines:
        assert any(p.text.endswith('\n') for p in doc.paragraphs if p.text.startswith('X'))


def test_trailing_source_whitespace_does_not_add_empty_body_paragraph(tmp_path, face):
    def draw(page):
        page.insert_text((30,60), 'A last paragraph.', fontname='fixture', fontsize=10)
        page.insert_text((30,72), '   ', fontname='fixture', fontsize=10)
    info, mapping = source(tmp_path, face, draw)
    target = tmp_path/'no-empty-tail.docx'
    convert_pdf(info, target, mapping, mode='editable', keep_line_breaks=True)
    assert Document(target).paragraphs[-1].text == 'A last paragraph.'


def test_retained_lines_round_available_page_width_outward(tmp_path, face):
    def draw(page):
        page.insert_text((8.023,60), 'A full source row '*5, fontname='fixture', fontsize=10)
    info, mapping = source(tmp_path, face, draw)
    target = tmp_path/'margin.docx'
    convert_pdf(info, target, mapping, mode='editable', keep_line_breaks=True)
    with pymupdf.open(info.path) as pdf:
        bounds = pdf[0].get_text('blocks')[0][:4]
    sec = Document(target).sections[0]
    assert sec.left_margin.pt <= bounds[0]
    assert sec.page_width.pt-sec.right_margin.pt >= bounds[2]+.075


@pytest.mark.parametrize('keep_lines', [False, True])
@pytest.mark.parametrize('wide_x', [30, 112])
def test_heading_widening_below_logo_keeps_reading_order(tmp_path, face, keep_lines, wide_x):
    bold = read_font(Path('C:/Windows/Fonts/arialbd.ttf'), installed=True)
    path = tmp_path/'wide-heading.pdf'
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=500, height=700)
        page.insert_font(fontname='regular', fontfile=face.path)
        page.insert_font(fontname='bold', fontfile=bold.path)
        page.draw_circle((90,40),20,color=(0,0,0),fill=(0,0,0))
        # Second row is below the logo but extends left of its column. Third
        # row falls within the old logo-neighbor search window again.
        for x,y,text in [(280,40,'FIRST HEADING'),(wide_x,73,'SECOND WIDE HEADING'),(280,85,'THIRD HEADING')]:
            page.insert_text((x,y),text,fontname='bold',fontsize=10)
        page.insert_text((30,120),'Body text follows the complete heading.',fontname='regular',fontsize=10)
        pdf.save(path)
    info=inspect_pdf(path)
    mapping={name:bold if 'Bold' in name else face for name in info.fonts}
    target=tmp_path/'wide-heading.docx'
    result=convert_pdf(info,target,mapping,mode='editable',keep_line_breaks=keep_lines)
    assert result['content_check']=='passed'
    doc=Document(target)
    text=''.join(t.text or '' for t in doc.element.iter(qn('w:t')))
    assert text.index('FIRST HEADING') < text.index('SECOND WIDE HEADING') < text.index('THIRD HEADING') < text.index('Body text')
    assert len(doc.tables)==1
    assert 'FIRST HEADING' in doc.tables[0].cell(0,1).text
    assert 'THIRD HEADING' not in doc.tables[0].cell(0,1).text
    assert any('SECOND WIDE HEADING' in p.text for p in doc.paragraphs)
    assert not list(doc.element.iter(qn('w:txbxContent')))
    assert not list(doc.element.iter(qn('w:framePr')))
    assert not list(doc.element.iter(qn('w:trHeight')))


def test_right_logo_shares_header_with_heading_and_subtitle(tmp_path, face):
    bold = read_font(Path('C:/Windows/Fonts/arialbd.ttf'), installed=True)
    path = tmp_path/'header.pdf'
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=320, height=800)
        page.insert_font(fontname='regular', fontfile=face.path)
        page.insert_font(fontname='bold', fontfile=bold.path)
        page.insert_text((20,45), 'Receipt heading', fontname='bold', fontsize=18)
        page.insert_text((20,65), 'A subtitle', fontname='regular', fontsize=14)
        page.draw_circle((276,47),25,color=(0,0,0),fill=(0,0,0))
        page.insert_text((20,100), 'Ordinary body paragraph.', fontname='regular', fontsize=12)
        pdf.save(path)
    info = inspect_pdf(path)
    mapping = {name:bold if 'Bold' in name else face for name in info.fonts}
    target = tmp_path/'header.docx'
    convert_pdf(info, target, mapping, mode='editable', keep_line_breaks=True)
    doc = Document(target)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert 'Receipt heading' in table.cell(0,0).text and 'A subtitle' in table.cell(0,0).text
    assert len(list(table.cell(0,1)._tc.iter(qn('w:drawing')))) == 1
    assert table.columns[0].width.pt == pytest.approx(231,abs=.05)
    assert table.columns[1].width.pt == pytest.approx(50,abs=.05)
    assert any(p.text == 'Ordinary body paragraph.' for p in doc.paragraphs)
    assert not list(doc.element.iter(qn('w:framePr')))
    assert not list(doc.element.iter(qn('w:txbxContent')))
    assert not list(doc.element.iter(qn('w:trHeight')))


@pytest.mark.parametrize('boxed', [False, True])
@pytest.mark.parametrize('order', ['row', 'column'])
def test_columns_rejected_before_publication(tmp_path, face, boxed, order):
    def draw(page):
        positions = [(column, row) for column in range(2) for row in range(3)]
        if order == 'row':
            positions.sort(key=lambda x: (x[1], x[0]))
        for column, row in positions:
            page.insert_text((35+column*260, 80+row*20), f'Column {column} line {row}', fontname='fixture', fontsize=12)
        if boxed:
            page.draw_rect(pymupdf.Rect(25, 55, 560, 135))
    info, mapping = source(tmp_path, face, draw)
    with pytest.raises(UnsupportedLayout, match='columns'):
        convert_pdf(info, tmp_path/'result.docx', mapping, mode='editable')
    assert not (tmp_path/'result.docx').exists()
    assert not (tmp_path/'result.partial.docx').exists()


@pytest.mark.parametrize('shape', ['vertical', 'four-stroke-box', 'incomplete-grid'])
def test_unsupported_drawing_cannot_disappear(tmp_path, face, shape):
    def draw(page):
        page.insert_text((35, 50), 'Drawing test', fontname='fixture', fontsize=12)
        if shape == 'vertical':
            page.draw_line((30, 70), (30, 140))
        elif shape == 'four-stroke-box':
            for a, b in [((30, 70), (300, 70)), ((300, 70), (300, 140)), ((300, 140), (30, 140)), ((30, 140), (30, 70))]:
                page.draw_line(a, b)
        else:
            page.draw_rect(pymupdf.Rect(30, 70, 300, 150))
            page.draw_line((30, 110), (200, 110))
            page.insert_text((40, 95), 'Cell', fontname='fixture')
    info, mapping = source(tmp_path, face, draw)
    with pytest.raises(UnsupportedLayout):
        convert_pdf(info, tmp_path/'result.docx', mapping, mode='editable')
    assert not (tmp_path/'result.docx').exists()


def test_cyrillic_wrapped_lines_keep_word_boundary_and_inline_fonts(tmp_path, face):
    bold = read_font(Path('C:/Windows/Fonts/arialbd.ttf'), installed=True)
    def draw(page):
        page.insert_font(fontname='strong', fontfile=bold.path)
        page.insert_text((30, 60), 'Начало предложения и', fontname='fixture', fontsize=12)
        page.insert_text((30, 74), 'продолжение строки.', fontname='fixture', fontsize=12)
        page.insert_text((30, 110), 'Жирный', fontname='strong', fontsize=12)
        x = 30+pymupdf.Font(fontfile=bold.path).text_length('Жирный', fontsize=12)
        page.insert_text((x, 110), 'Текст', fontname='fixture', fontsize=12)
    info, mapping = source(tmp_path, face, draw)
    mapping = {n: bold if 'Bold' in n else face for n in info.fonts}
    with pymupdf.open(info.path) as pdf:
        doc, expected, _ = build_document(pdf, mapping)
    verify_document(doc, expected)
    texts = '\n'.join(p.text for p in doc.paragraphs)
    assert 'и продолжение' in texts
    assert 'ЖирныйТекст' in texts
    assert any(r.bold and r.text == 'Жирный' for p in doc.paragraphs for r in p.runs)
    target_run = next(r for p in doc.paragraphs for r in p.runs if r.text == 'Жирный')
    target_run.bold = False
    with pytest.raises(ValueError, match='emphasis'):
        verify_document(doc, expected)


def test_cancel_and_existing_partial_are_preserved(tmp_path, face):
    info, mapping = source(tmp_path, face)
    target = tmp_path/'result.docx'
    cancel = threading.Event()
    cancel.set()
    with pytest.raises(InterruptedError):
        convert_pdf(info, target, mapping, mode='editable', cancel=cancel)
    assert not target.exists()
    partial = tmp_path/'result.partial.docx'
    partial.write_bytes(b'user-owned')
    with pytest.raises(FileExistsError):
        convert_pdf(info, target, mapping, mode='editable')
    assert partial.read_bytes() == b'user-owned'
    assert not target.exists()


def test_editable_batch_reuses_folder_without_reports(tmp_path, face):
    info, mapping = source(tmp_path, face)
    out = tmp_path/'out'
    first = convert_batch([info], out, mapping, mode='editable')
    original = (out/'form.docx').read_bytes()
    second = convert_batch([info], out, mapping, mode='editable')
    assert first[0]['mode'] == second[0]['mode'] == 'editable'
    assert second[0]['output'] == 'form (2).docx'
    assert (out/'form.docx').read_bytes() == original
    assert {p.suffix for p in out.iterdir()} == {'.docx'}


def test_unknown_mode_does_not_create_output(tmp_path, face):
    info, mapping = source(tmp_path, face)
    with pytest.raises(ValueError, match='Choose Editable'):
        convert_batch([info], tmp_path/'out', mapping, mode='wrong')
    assert not (tmp_path/'out').exists()


def test_narrative_parenthetical_is_not_reordered_as_field_caption(tmp_path, face):
    def draw(page):
        page.insert_text((30, 50), 'Name: _____________________', fontname='fixture', fontsize=12)
        page.insert_text((30, 100), 'Please read these instructions before you fill in the form.', fontname='fixture', fontsize=12)
        page.insert_text((30, 114), '(This instruction applies to all applicants.)', fontname='fixture', fontsize=12)
    info, mapping = source(tmp_path, face, draw)
    target = tmp_path/'result.docx'
    convert_pdf(info, target, mapping, mode='editable')
    doc = Document(target)
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert text.index('Please read') < text.index('(This instruction')
    assert len(doc.tables) == 1


def test_immediate_field_caption_has_its_own_following_row(tmp_path, face):
    def draw(page):
        page.insert_text((30, 50), 'Name: ____________________________', fontname='fixture', fontsize=12)
        page.insert_text((80, 64), '(Full name)', fontname='fixture', fontsize=12)
        page.insert_text((30, 100), 'Following paragraph.', fontname='fixture', fontsize=12)
    info, mapping = source(tmp_path, face, draw)
    target = tmp_path/'result.docx'
    convert_pdf(info, target, mapping, mode='editable')
    doc = Document(target)
    assert len(doc.tables) == 2
    assert '(Full name)' in doc.tables[1].cell(0, 1).text.replace('\xa0', ' ')
    assert all(p.paragraph_format.keep_with_next for c in doc.tables[0].rows[0].cells for p in c.paragraphs)
    separator = doc.tables[0]._tbl.getnext()
    assert separator.find('./'+qn('w:pPr')+'/'+qn('w:keepNext')) is not None
    assert any('Following paragraph.' in p.text for p in doc.paragraphs)


def test_signature_label_stays_with_following_field(tmp_path, face):
    def draw(page):
        page.insert_text((30, 60), 'Signature of applicant', fontname='fixture', fontsize=12)
        page.insert_text((30, 85), '____________________________', fontname='fixture', fontsize=12)
    info, mapping = source(tmp_path, face, draw)
    with pymupdf.open(info.path) as pdf:
        doc, expected, _ = build_document(pdf, mapping)
    verify_document(doc, expected)
    label = next(p for p in doc.paragraphs if 'Signature of applicant' in p.text)
    assert label.paragraph_format.keep_with_next


@pytest.mark.parametrize('order', ['row', 'column'])
def test_ruled_table_cell_order_and_expanding_rows(tmp_path, face, order):
    def draw(page):
        for x in [30, 220, 500]:
            page.draw_line((x, 70), (x, 170))
        for y in [70, 120, 170]:
            page.draw_line((30, y), (500, y))
        positions = [(r,c,k) for r in range(2) for c in range(2) for k in range(2)]
        if order == 'column':
            positions.sort(key=lambda p:(p[1],p[0],p[2]))
        for r,c,k in positions:
            page.insert_text((40+c*190, 90+r*50+k*14), f'cell {r} {c} line {k}', fontname='fixture', fontsize=10)
        page.insert_text((30,220), 'Following paragraph.', fontname='fixture', fontsize=10)
    info,mapping=source(tmp_path,face,draw)
    target=tmp_path/'result.docx'; convert_pdf(info,target,mapping,mode='editable')
    doc=Document(target)
    assert len(doc.tables) == 1
    table=doc.tables[0]
    assert len(table.rows) == len(table.columns) == 2
    for r in range(2):
        for c in range(2):
            assert f'cell {r} {c} line 0' in table.cell(r,c).text
            assert f'cell {r} {c} line 1' in table.cell(r,c).text
    assert not list(doc.element.iter(qn('w:trHeight')))
    assert not list(doc.element.iter(qn('w:txbxContent')))
    assert len(list(doc.element.iter(qn('w:tcBorders')))) == 4
    for pr in doc.element.iter(qn('w:tcPr')):
        tags=[n.tag for n in pr]
        if qn('w:tcBorders') in tags and qn('w:tcMar') in tags:
            assert tags.index(qn('w:tcBorders')) < tags.index(qn('w:tcMar'))
    table.cell(0,1).add_paragraph('Additional editable text\nAnother line')
    doc.save(target)
    assert 'Additional' in Document(target).tables[0].cell(0,1).text


def test_text_crossing_table_border_is_rejected(tmp_path, face):
    def draw(page):
        page.draw_rect(pymupdf.Rect(30,70,400,170))
        page.draw_line((30,120),(400,120))
        page.draw_line((150,70),(150,170))
        page.insert_text((130,100),'Text crosses a border',fontname='fixture',fontsize=12)
    info,mapping=source(tmp_path,face,draw)
    with pytest.raises(UnsupportedLayout,match='crosses'):
        convert_pdf(info,tmp_path/'result.docx',mapping,mode='editable')


@pytest.mark.parametrize('edge', ['top', 'bottom'])
@pytest.mark.parametrize('keep_lines', [False, True])
def test_outside_blank_grazing_form_border_remains_editable(tmp_path, face, edge, keep_lines):
    def draw(page):
        page.insert_text((40,100),'____________________',fontname='fixture',fontsize=12)
        bounds = pymupdf.Rect(page.get_text('blocks')[0][:4])
        if edge == 'top':
            box = pymupdf.Rect(30,bounds.y1-.7,400,190)
            y = 150
        else:
            box = pymupdf.Rect(30,20,400,bounds.y0+.7)
            y = 50
        page.draw_rect(box)
        page.insert_text((40,y),'Inside the form.',fontname='fixture',fontsize=12)
    info,mapping=source(tmp_path,face,draw)
    target=tmp_path/'touching.docx'
    result=convert_pdf(info,target,mapping,mode='editable',keep_line_breaks=keep_lines)
    doc=Document(target)
    assert result['fields']==1
    assert len(doc.tables)==2
    assert any('Inside the form.' in t.cell(0,0).text for t in doc.tables)
    assert not list(doc.element.iter(qn('w:trHeight')))


@pytest.mark.parametrize('text,overlap,x', [('Label ____',.7,40),('____________',3,40),('____________',.7,25)])
def test_real_text_deep_or_side_crossing_form_border_still_rejected(tmp_path,face,text,overlap,x):
    def draw(page):
        page.insert_text((x,100),text,fontname='fixture',fontsize=12)
        bounds=pymupdf.Rect(page.get_text('blocks')[0][:4])
        page.draw_rect(pymupdf.Rect(30,bounds.y1-overlap,400,190))
        page.insert_text((40,150),'Inside.',fontname='fixture',fontsize=12)
    info,mapping=source(tmp_path,face,draw)
    with pytest.raises(UnsupportedLayout,match='crosses a form border'):
        convert_pdf(info,tmp_path/'crossing.docx',mapping,mode='editable')


def test_indented_text_is_not_centered_and_line_spaces_are_normalized(tmp_path, face):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    def draw(page):
        page.insert_text((30,50),'Main text sets the margin.',fontname='fixture',fontsize=12)
        page.insert_text((120,100),'Indented paragraph with ',fontname='fixture',fontsize=12)
        page.insert_text((120,114),' continuation of the same sentence.',fontname='fixture',fontsize=12)
    info,mapping=source(tmp_path,face,draw)
    target=tmp_path/'result.docx';convert_pdf(info,target,mapping,mode='editable')
    doc=Document(target)
    para=next(p for p in doc.paragraphs if p.text.startswith('Indented'))
    assert para.alignment != WD_ALIGN_PARAGRAPH.CENTER
    assert para.paragraph_format.left_indent.pt == pytest.approx(90,abs=1)
    assert 'with continuation' in para.text and '  ' not in para.text


def test_unrepresentable_grid_border_is_rejected(tmp_path, face):
    def draw(page):
        page.draw_rect(pymupdf.Rect(30,70,400,170),width=16)
        page.draw_line((30,120),(400,120),width=16)
        page.insert_text((50,100),'Cell',fontname='fixture',fontsize=12)
    info,mapping=source(tmp_path,face,draw)
    with pytest.raises(UnsupportedLayout):
        convert_pdf(info,tmp_path/'result.docx',mapping,mode='editable')


@pytest.mark.parametrize('mode', ['editable', 'layout'])
def test_conversion_embeds_only_used_faces(tmp_path, face, mode):
    from pagethaw.font_embedding import verify_embedded_fonts
    info,mapping=source(tmp_path,face)
    mapping['unused']=read_font(Path('C:/Windows/Fonts/arialbd.ttf'),installed=True)
    target=tmp_path/'embedded.docx'
    result=convert_pdf(info,target,mapping,mode=mode,embed=True)
    assert len(result['embedded_fonts'])==1 and not result['font_embedding_skipped']
    verify_embedded_fonts(target,result)
