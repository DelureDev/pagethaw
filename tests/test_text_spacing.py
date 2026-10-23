from dataclasses import replace

import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from fontTools.ttLib import TTFont

from pagethaw.editable import build_document, lines_on, verify_document
from pagethaw.fonts import read_font
from pagethaw.text_spacing import TextSpacing
from test_font_embedding import make_font


@pytest.fixture
def tall_font(tmp_path):
    path = tmp_path/'tall.ttf'
    make_font(path)
    with TTFont(path) as font:
        font['OS/2'].version = 4
        font['OS/2'].fsSelection |= 128
        font['OS/2'].sTypoAscender = 1000
        font['OS/2'].sTypoDescender = -300
        font['OS/2'].sTypoLineGap = 20
        font.save(path)
    return read_font(path)


def pdf_lines(face, baselines, texts=None, sizes=None):
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_font(fontname='test', fontfile=face.path)
    for i, y in enumerate(baselines):
        page.insert_text((30, y), (texts or ['ordinary flowing text']*len(baselines))[i],
                         fontname='test', fontsize=(sizes or [10]*len(baselines))[i])
    return pdf, lines_on(page)


def test_multiline_keeps_sizes_with_responsive_leading(tall_font, tmp_path):
    with pdf_lines(tall_font, [60, 72, 84])[0] as pdf:
        mapping = {s['font']:tall_font for l in lines_on(pdf[0]) for s in l.spans}
        doc, expected, _ = build_document(pdf, mapping)
    verify_document(doc, expected)
    target = tmp_path/'flow.docx'
    doc.save(target)
    saved = Document(target)
    para = next(p for p in saved.paragraphs if len(p.runs) > 1)
    assert para.paragraph_format.line_spacing_rule == WD_LINE_SPACING.MULTIPLE
    # Word's line multiple compensates the tall font metrics, not the type size.
    assert 12 <= para.paragraph_format.line_spacing*13.2 <= 12.06
    assert all(r.font.size.pt == 10 for r in para.runs)
    assert not list(saved.element.iter(qn('w:trHeight')))


def test_legacy_font_and_missing_metrics_keep_expanding_fallback(tmp_path):
    face = make_font(tmp_path/'legacy.ttf')
    with pdf_lines(face, [60, 72, 84])[0] as pdf:
        lines = lines_on(pdf[0]); mapping = {s['font']:face for l in lines for s in l.spans}
        doc, _, _ = build_document(pdf, mapping)
        para = next(p for p in doc.paragraphs if p.text)
        assert para.paragraph_format.line_spacing_rule == WD_LINE_SPACING.AT_LEAST
        missing = {name:replace(f, path=str(tmp_path/'absent.ttf')) for name,f in mapping.items()}
        assert TextSpacing(missing).multiple(lines) is None


@pytest.mark.parametrize('baselines,sizes', [([60,72,87],[10,10,10]), ([60,72,84],[10,11,10])])
def test_mixed_sizes_or_irregular_multiline_spacing_not_compressed(tall_font, baselines, sizes):
    pdf, lines = pdf_lines(tall_font, baselines, sizes=sizes)
    try:
        spacing = TextSpacing({s['font']:tall_font for l in lines for s in l.spans})
        assert spacing.multiple(lines) is None
    finally:
        pdf.close()


def test_single_rows_need_repeated_local_evidence_and_reset_between_containers(tall_font):
    pdf, lines = pdf_lines(tall_font, [60,72,84,130])
    try:
        spacing = TextSpacing({s['font']:tall_font for l in lines for s in l.spans})
        spacing.flow(lines)
        assert spacing.multiple([lines[1]]) is not None
        assert spacing.multiple([lines[3]]) is None  # not the page's body default
        spacing.flow(lines[:2])
        assert spacing.multiple([lines[0]]) is None  # one pair is insufficient
        spacing.flow(lines[3:])
        assert spacing.multiple([lines[1]]) is None  # no stale container evidence
    finally:
        pdf.close()


def test_form_cells_grow_and_empty_caption_marks_have_caption_font(tall_font, tmp_path):
    texts = ['Name: __________ Date: _________', '(name)              (date)', 'ordinary text.']
    with pdf_lines(tall_font, [60,72,84], texts)[0] as pdf:
        mapping = {s['font']:tall_font for l in lines_on(pdf[0]) for s in l.spans}
        doc, expected, _ = build_document(pdf, mapping)
    verify_document(doc, expected)
    caption = doc.tables[1]
    for cell in caption.rows[0].cells:
        mark = cell.paragraphs[0]._p.find('./'+qn('w:pPr')+'/'+qn('w:rPr'))
        assert mark.find(qn('w:rFonts')).get(qn('w:ascii')) == tall_font.family
        assert mark.find(qn('w:sz')).get(qn('w:val')) == '20'
    field = doc.tables[0].cell(0,1).paragraphs[0]
    assert field.paragraph_format.line_spacing_rule == WD_LINE_SPACING.MULTIPLE
    field.runs[0].text = 'First\nSecond\nThird'
    doc.save(tmp_path/'edited.docx')
    saved = Document(tmp_path/'edited.docx')
    assert 'Third' in saved.tables[0].cell(0,1).text
    assert not list(saved.element.iter(qn('w:trHeight')))
    assert not list(saved.element.iter(qn('w:txbxContent')))
    for p in saved.paragraphs:
        if p.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY:
            assert not p.text.strip()  # fixed tiny spacing only separates tables
            assert p.paragraph_format.line_spacing.twips == 1
