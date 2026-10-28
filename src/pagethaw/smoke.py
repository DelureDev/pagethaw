"""Synthetic test that can also run inside the frozen portable application."""
from pathlib import Path
import json
import sys
import os
import re
import pymupdf

from .engine import inspect_pdf, convert_batch
from .fonts import installed_fonts


def run(folder, compare=False):
    folder = Path(folder).resolve()
    folder.mkdir(parents=True, exist_ok=False)
    faces = installed_fonts()
    face = next((f for f in faces if f.family == 'Arial' and not f.bold and not f.italic), None)
    if face is None:
        raise RuntimeError('The synthetic smoke test needs the Windows Arial regular font.')
    source = folder/'synthetic-form.pdf'
    with pymupdf.open() as pdf:
        for page_number in range(2):
            page = pdf.new_page()
            page.insert_font(fontname='testfont', fontfile=face.path)
            page.insert_text((35, 55), f'PageThaw sample form — page {page_number+1}', fontname='testfont', fontsize=16)
            page.insert_text((35, 90), 'Name: _____________________', fontname='testfont', fontsize=11)
            page.insert_text((35, 115), 'Reference: account_id / Test 123', fontname='testfont', fontsize=11)
            page.draw_rect(pymupdf.Rect(30, 140, 560, 240), color=(.1, .5, .55))
            page.insert_text((40, 165), 'An editable form with a retained graphic border.', fontname='testfont', fontsize=11)
            page.insert_text((35, 265), 'This paragraph continues', fontname='testfont', fontsize=11)
            page.insert_text((35, 279), 'as ordinary editable text.', fontname='testfont', fontsize=11)
        pdf.save(source)
    info = inspect_pdf(source)
    assert not info.error, info.error
    output = folder/'output'
    records = convert_batch([info], output, {name: face for name in info.fonts})
    assert records[0].get('content_check') == 'passed', records
    original = (output/records[0]['output']).read_bytes()
    repeated = convert_batch([info], output, {name: face for name in info.fonts}, rules=True)
    assert repeated[0].get('content_check') == 'passed', repeated
    assert repeated[0]['output'] == 'synthetic-form (2).docx'
    assert (output/records[0]['output']).read_bytes() == original
    assert {p.name for p in output.iterdir()} == {'synthetic-form.docx', 'synthetic-form (2).docx'}
    editable = convert_batch([info], folder/'editable', {name: face for name in info.fonts}, mode='editable', embed=True)
    assert editable[0].get('content_check') == 'passed', editable
    from docx import Document
    from docx.oxml.ns import qn
    editable_doc = Document(folder/'editable'/editable[0]['output'])
    assert not list(editable_doc.element.iter(qn('w:txbxContent')))
    assert editable_doc.tables and editable_doc.paragraphs
    assert editable[0]['embedded_fonts'] and not editable[0]['font_embedding_skipped'], editable
    retained = convert_batch([info], folder/'retained-lines', {name: face for name in info.fonts},
                             mode='editable', keep_line_breaks=True)
    assert retained[0].get('content_check') == 'passed', retained
    retained_doc = Document(folder/'retained-lines'/retained[0]['output'])
    assert any('This paragraph continues\nas ordinary editable text.' in p.text for p in retained_doc.paragraphs)
    assert not list(retained_doc.element.iter(qn('w:txbxContent')))
    assert not list(retained_doc.element.iter(qn('w:framePr')))
    grid_source = folder/'synthetic-table.pdf'
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_font(fontname='testfont', fontfile=face.path)
        for x in (30, 250, 550):
            page.draw_line((x, 60), (x, 160))
        for y in (60, 110, 160):
            page.draw_line((30, y), (550, y))
        for r in range(2):
            for c in range(2):
                page.insert_text((40+c*220, 85+r*50), f'Cell {r} {c}', fontname='testfont', fontsize=12)
        pdf.save(grid_source)
    grid_info = inspect_pdf(grid_source)
    grid_records = convert_batch([grid_info], folder/'table', {name: face for name in grid_info.fonts}, mode='editable')
    assert grid_records[0].get('content_check') == 'passed', grid_records
    table = Document(folder/'table'/grid_records[0]['output']).tables[0]
    assert len(table.rows) == len(table.columns) == 2 and 'Cell 1 1' in table.cell(1, 1).text
    # Exercise both fixes in the actual frozen executable on a clean Windows VM.
    notes_source = folder/'synthetic-notes.pdf'
    font = pymupdf.Font(fontfile=face.path)
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=500, height=500)
        page.insert_font(fontname='testfont', fontfile=face.path)
        prefix = 'Body with a footnote'
        page.insert_text((30, 50), prefix, fontname='testfont', fontsize=12)
        page.insert_text((30+font.text_length(prefix, fontsize=12), 46), '7', fontname='testfont', fontsize=7)
        page.insert_text((30, 140), '_'*50, fontname='testfont', fontsize=12)
        blank = next(line for block in page.get_text('dict')['blocks'] for line in block.get('lines', [])
                     if ''.join(s['text'] for s in line['spans']).startswith('___'))
        page.draw_rect(pymupdf.Rect(25, blank['bbox'][3]-.8, 475, 220))
        page.insert_text((30, 180), 'Inside the form.', fontname='testfont', fontsize=12)
        page.draw_line((30, 390), (300, 390))
        page.insert_text((30, 406), '7', fontname='testfont', fontsize=7)
        page.insert_text((30+font.text_length('7', fontsize=7), 410), ' Editable note content.', fontname='testfont', fontsize=12)
        page = pdf.new_page(width=500, height=500)
        page.insert_font(fontname='testfont', fontfile=face.path)
        page.insert_text((30, 50), 'Following source page.', fontname='testfont', fontsize=12)
        pdf.save(notes_source)
    notes_info = inspect_pdf(notes_source)
    notes_records = convert_batch([notes_info], folder/'notes', {name: face for name in notes_info.fonts}, mode='editable')
    assert notes_records[0].get('content_check') == 'passed', notes_records
    notes_doc = Document(folder/'notes'/notes_records[0]['output'])
    from .footnotes import read_part
    note_text = ''.join(t.text or '' for t in read_part(notes_doc).iter(qn('w:t')))
    assert note_text == '7 Editable note content.'
    assert len(list(notes_doc.element.iter(qn('w:footnoteReference')))) == 1
    assert 'Editable note content.' not in ''.join(p.text for p in notes_doc.paragraphs)
    continuation = next(p for p in notes_doc.paragraphs if 'Following source page.' in p.text)
    assert continuation._p.getprevious().xpath('.//w:br[@w:type="page"]')
    assert not list(notes_doc.element.iter(qn('w:txbxContent')))
    assert not list(notes_doc.element.iter(qn('w:framePr')))
    fields_source = folder/'synthetic-fields.pdf'
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_font(fontname='testfont', fontfile=face.path)
        row = 'Name: __________ Date «____» ______ 20____y.'
        page.insert_text((30, 60), row, fontname='testfont', fontsize=10)
        page.insert_text((30, 90), 'Ordinary text: account_id 20____.', fontname='testfont', fontsize=10)
        pdf.save(fields_source)
    fields_info = inspect_pdf(fields_source)
    fields_records = convert_batch([fields_info], folder/'fields', {n: face for n in fields_info.fonts}, mode='editable')
    assert fields_records[0].get('content_check') == 'passed' and fields_records[0]['fields'] == 4, fields_records
    fields_doc = Document(folder/'fields'/fields_records[0]['output'])
    field_table = fields_doc.tables[0]
    with pymupdf.open(fields_source) as pdf:
        raw_line = pdf[0].get_text('rawdict')['blocks'][0]['lines'][0]
        chars = [c for span in raw_line['spans'] for c in span['chars']]
    expected_edges = [(chars[m.start()]['bbox'][0]-30, chars[m.end()-1]['bbox'][2]-30)
                      for m in re.finditer(r'_{3,}', row)]
    actual_edges, edge = [], 0
    for column, cell in zip(field_table.columns, field_table.rows[0].cells):
        if cell._tc.xpath('./w:tcPr/w:tcBorders/w:bottom[@w:val="single"]'):
            actual_edges.append((edge, edge+column.width.pt))
        edge += column.width.pt
    assert len(actual_edges) == len(expected_edges) == 4
    assert all(abs(a-e) <= .2 for actual, expected in zip(actual_edges, expected_edges)
               for a, e in zip(actual, expected)), (actual_edges, expected_edges)
    assert field_table.rows[0].cells[-2]._tc.xpath('./w:tcPr/w:tcBorders/w:bottom[@w:val="single"]')
    assert field_table.rows[0].cells[-1].text == 'y.'
    assert not list(fields_doc.element.iter(qn('w:trHeight')))
    assert 'Ordinary text: account_id 20____.' in '\n'.join(p.text for p in fields_doc.paragraphs)
    continued_source = folder/'synthetic-continued.pdf'
    with pymupdf.open() as pdf:
        for number in range(2):
            page = pdf.new_page(width=500, height=500)
            page.insert_font(fontname='testfont', fontfile=face.path)
            page.insert_text((30,65), 'Yes        / No        agree: __________', fontname='testfont', fontsize=10)
            raw = page.get_text('rawdict')['blocks'][0]['lines'][0]['spans'][0]['chars']
            x = raw[4]['bbox'][0]
            page.draw_rect(pymupdf.Rect(x,58,x+7,65))
            points = ([(20,490),(20,40),(470,40),(470,490)] if number == 0
                      else [(20,10),(20,100),(470,100),(470,10)])
            page.draw_polyline(points, closePath=False, width=1.5)
        pdf.save(continued_source)
    continued_info = inspect_pdf(continued_source)
    continued_records = convert_batch([continued_info], folder/'continued',
                                      {n: face for n in continued_info.fonts}, mode='editable')
    assert continued_records[0].get('content_check') == 'passed', continued_records
    continued_doc = Document(folder/'continued'/continued_records[0]['output'])
    assert len(continued_doc.tables) == 1 and len(continued_doc.tables[0].rows) > 1
    assert len(continued_doc.element.xpath('.//wp:inline')) == 2
    assert not continued_doc.element.xpath('.//wp:anchor')
    assert not list(continued_doc.element.iter(qn('w:trHeight')))
    continued_borders = continued_doc.tables[0]._tbl.tblPr.find(qn('w:tblBorders'))
    assert continued_borders.find(qn('w:top')).get(qn('w:val')) == 'nil'
    assert continued_borders.find(qn('w:bottom')).get(qn('w:val')) == 'nil'
    assert continued_borders.find(qn('w:left')).get(qn('w:sz')) == '12'
    highlight_source = folder/'synthetic-highlight.pdf'
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_font(fontname='testfont', fontfile=face.path)
        page.insert_text((30,60), 'Highlighted editable text', fontname='testfont', fontsize=10)
        rect = page.search_for('Highlighted')[0]
        annot = page.add_highlight_annot(rect)
        annot.set_colors(stroke=(1,.5,0))
        annot.set_opacity(.4)
        annot.set_blendmode('Multiply')
        annot.update()
        xref = int(pdf.xref_get_key(annot.xref, 'AP/N')[1].split()[0])
        pdf.xref_set_key(xref, 'Resources', '<< /ExtGState << /H << /ca .4 /BM /Multiply >> >> >>')
        pdf.update_stream(xref, f'q /H gs 1 .5 0 rg {rect.x0} {page.rect.height-rect.y1} {rect.width} {rect.height} re f Q'.encode())
        pdf.save(highlight_source)
    highlight_info = inspect_pdf(highlight_source)
    highlight_records = convert_batch([highlight_info], folder/'highlight',
                                     {n: face for n in highlight_info.fonts}, mode='editable')
    assert highlight_records[0].get('content_check') == 'passed', highlight_records
    highlight_doc = Document(folder/'highlight'/highlight_records[0]['output'])
    shaded_runs = highlight_doc.element.xpath('.//w:r[w:rPr/w:shd[@w:fill="FFCC99"]]')
    assert ''.join(t.text for r in shaded_runs for t in r.iter(qn('w:t'))) == 'Highlighted'
    assert not list(highlight_doc.element.iter(qn('w:txbxContent')))
    bold_face = next(f for f in faces if f.family == 'Arial' and f.bold and not f.italic)
    header_source = folder/'synthetic-heading.pdf'
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=500, height=700)
        page.insert_font(fontname='regular', fontfile=face.path)
        page.insert_font(fontname='bold', fontfile=bold_face.path)
        page.draw_circle((90,40),20,color=(0,0,0),fill=(0,0,0))
        for x,y,text in [(280,40,'FIRST HEADING'),(112,73,'SECOND WIDE HEADING'),(280,85,'THIRD HEADING')]:
            page.insert_text((x,y),text,fontname='bold',fontsize=10)
        page.insert_text((30,120),'Body follows the heading.',fontname='regular',fontsize=10)
        pdf.save(header_source)
    header_info = inspect_pdf(header_source)
    header_records = convert_batch([header_info], folder/'heading',
                                  {n: bold_face if 'Bold' in n else face for n in header_info.fonts}, mode='editable')
    assert header_records[0].get('content_check') == 'passed', header_records
    header_doc = Document(folder/'heading'/header_records[0]['output'])
    header_text = ''.join(t.text or '' for t in header_doc.element.iter(qn('w:t')))
    assert header_text.index('FIRST HEADING') < header_text.index('SECOND WIDE HEADING') < header_text.index('THIRD HEADING') < header_text.index('Body follows')
    assert len(header_doc.tables) == 1 and 'THIRD HEADING' not in header_doc.tables[0].cell(0,1).text
    comparison = 'not run'
    if compare:
        from .word_review import review_batch
        review_batch([info], output, records=repeated, review_output=folder/'review')
        results = json.loads((folder/'review/comparison-report.json').read_text(encoding='utf-8'))
        assert results[0].get('page_count_check') == 'passed', results
        assert len(list((folder/'review/comparisons/1').glob('page-*.png'))) == 2
        comparison = 'exported; visual inspection required'
    from . import gui
    old_local = os.environ.get('LOCALAPPDATA')
    os.environ['LOCALAPPDATA'] = str(folder/'test-profile')
    app = None
    try:
        app = gui.App()
        app.withdraw()
        app.language.set('en')
        app.switch_language()
        assert app.language.get() == 'ru'
        assert app.font_frame.winfo_manager() == 'grid'
        app.switch_language()
        assert app.language.get() == 'en'
        assert app.font_frame.winfo_manager() == 'grid'
        assert app.mode.get() == 'editable'
        assert 'Normal paragraphs' in app.mode_help.cget('text')
        assert app.keep_lines.get() is False
        app.keep_lines.set(True)
        app.mode_changed()
        app.switch_language()
        assert app.keep_lines.get() and app.keep_lines_option.cget('text') == app.t('Keep PDF line breaks in normal paragraphs')
        app.mode.set('layout')
        app.mode_changed()
        assert 'disabled' in app.keep_lines_option.state()
        app.update_idletasks()
    finally:
        if app:
            app.close()
        if old_local is None:
            os.environ.pop('LOCALAPPDATA', None)
        else:
            os.environ['LOCALAPPDATA'] = old_local
    bundled = 'not applicable'
    if getattr(sys, 'frozen', False):
        root = Path(sys._MEIPASS).resolve()
        for name in ('pymupdf', 'docx', 'fontTools', 'PIL', 'tkinter'):
            module = __import__(name)
            assert Path(module.__file__).resolve().is_relative_to(root), name
        bundled = 'passed'
    (folder/'smoke-result.json').write_text(json.dumps({'status': 'passed', 'pages': 2, 'records': records,
        'frozen': bool(getattr(sys, 'frozen', False)), 'bundled_runtime': bundled, 'gui': 'passed',
        'repeat_conversion': 'passed', 'output_only_docx': 'passed', 'editable_mode': 'passed',
        'font_embedding': 'passed', 'ruled_tables': 'passed', 'retained_lines': 'passed',
        'native_footnotes': 'passed', 'touching_blank_border': 'passed', 'field_geometry': 'passed',
        'inline_squares_and_continuations': 'passed',
        'highlights_and_border_styles': 'passed',
        'contiguous_heading_order': 'passed',
        'comparison': comparison}, indent=2), encoding='utf-8')
    return records
