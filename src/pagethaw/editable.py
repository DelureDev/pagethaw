"""Flowing Word output for single-column forms, without positioned text boxes.

This deliberately supports less than the layout renderer. PDF line/span provenance
is kept until emission; ambiguous columns, grids and text/graphic overlays fail.
No PDF-specific names, wording or fonts are used to reconstruct the document.
"""
from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
import io
import math
import re
from statistics import median

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from .inline_forms import attach_square, continuation_borders
from .highlights import prepared_pdf, apply_highlights
from .ruled_tables import find_grids, partition
from .text_spacing import TextSpacing, source_tracking
from .footnotes import find_notes, create_part, add_note, read_part


class UnsupportedLayout(ValueError):
    pass


def reject(detail):
    raise UnsupportedLayout('Editable mode: '+detail+' Use Preserve layout for this PDF.')


def blank_ranges(text):
    """Standalone blanks and the unfinished year following day/month blanks.

    Numeric/letter identifiers stay literal. A year can touch a one-letter unit
    abbreviation, e.g. 20____y. Return only underscores, preserving the prefix
    and suffix as ordinary text with their original emphasis.
    """
    blanks = [match.span() for match in re.finditer(r'(?<!\w)_{3,}(?!\w)', text)]
    years = []
    for match in re.finditer(r'(?<!\w)\d{2}(_{3,})(?=\W|$|[^\W\d_]\.)', text):
        preceding = [(a, b) for a, b in blanks if b <= match.start()]
        if len(preceding) < 2:
            continue
        day, month = preceding[-2:]
        separators = text[day[1]:month[0]]+text[month[1]:match.start()]
        if not any(c.isalnum() or c == '_' for c in separators):
            years.append(match.span(1))
    return sorted(blanks+years)


@dataclass
class Line:
    spans: list
    bbox: pymupdf.Rect
    block: int

    @property
    def text(self):
        return ''.join(c['c'] for s in self.spans for c in s['chars'])

    @property
    def size(self):
        return max(s['size'] for s in self.spans)

    @property
    def field(self):
        return bool(blank_ranges(self.text))

    @property
    def bold(self):
        return all(s['flags'] & 16 for s in self.spans if any(c['c'].strip() for c in s['chars']))


def cancelled(event):
    if event and event.is_set():
        raise InterruptedError('Conversion cancelled.')


def lines_on(page):
    lines = []
    for bi, block in enumerate(page.get_text('rawdict')['blocks']):
        for raw in block.get('lines', []):
            spans = [s for s in raw['spans'] if s['chars']]
            if not spans:
                continue
            # PDF engines can expose horizontally separated portions of one row
            # as separate lines. Combine only the same block and baseline.
            y = spans[0]['origin'][1]
            same = next((l for l in reversed(lines) if l.block == bi and
                         abs(l.spans[0]['origin'][1]-y) < .5), None)
            if same:
                same.spans.extend(spans)
                same.spans.sort(key=lambda s: s['origin'][0])
                same.bbox |= pymupdf.Rect(raw['bbox'])
            else:
                lines.append(Line(spans, pymupdf.Rect(raw['bbox']), bi))
    return sorted(lines, key=lambda l: (round(l.bbox.y0, 1), l.bbox.x0))


def intersects(a, b, tolerance=.5):
    r = a & b
    return r.width > tolerance and r.height > tolerance


def contains(outer, inner, tolerance=1):
    return (outer+(-tolerance, -tolerance, tolerance, tolerance)).contains(inner)


def blank_touches_border(box, line):
    """An outside blank can have a descender bbox just across a horizontal edge.

    Preserve it as an ordinary expanding field. This exception never applies to
    labels, to vertical-edge crossings, or to a blank whose baseline is inside.
    """
    if not re.fullmatch(r'\s*_{3,}\s*', line.text):
        return False
    if line.bbox.x0 < box.x0 or line.bbox.x1 > box.x1:
        return False
    baselines = [s['origin'][1] for s in line.spans]
    overlap = min(box.y1, line.bbox.y1)-max(box.y0, line.bbox.y0)
    return (0 < overlap <= min(1.5, line.size*.2) and
            (max(baselines) <= box.y0 or min(baselines) >= box.y1))


def graphics_on(page, lines, continuations=None):
    """Separate simple outer boxes/rules from standalone graphics such as logos."""
    boxes, rules, rects = [], [], []
    drawings = page.get_drawings()
    grids = find_grids(drawings)
    handled = {index for g in grids for index in g.drawings}
    for index, draw in enumerate(drawings):
        if index in handled:
            continue
        if attach_square(page, draw, lines):
            continue
        if continuations and (page.number, index) in continuations:
            boxes.append(pymupdf.Rect(draw['rect']))
            continue
        rect = pymupdf.Rect(draw['rect'])
        items = draw['items']
        if draw['type'] == 's' and len(items) == 1 and items[0][0] == 're' and rect.width > 40 and rect.height > 25:
            boxes.append(rect)
        elif draw['type'] == 's' and rect.height < 1 and rect.width > 20:
            rules.append(rect)
        else:
            if rect.is_empty and (rect.width > .1 or rect.height > .1):
                reject('vertical or degenerate drawing paths are not supported.')
            rects.append(rect)
    rects.extend(pymupdf.Rect(i['bbox']) for i in page.get_image_info())
    for i, grid in enumerate(grids):
        if any(intersects(grid.bbox, other.bbox) for other in grids[i+1:]):
            reject('nested or overlapping table borders are not supported.')
        if any(grid.bbox.intersects(r+(-.1, -.1, .1, .1)) for r in boxes+rules+rects):
            reject('a ruled table contains unsupported graphics.')
        if any(min(l.bbox.y1, grid.bbox.y1)-max(l.bbox.y0, grid.bbox.y0) > 1 and
               (l.bbox.x1 < grid.bbox.x0 or l.bbox.x0 > grid.bbox.x1) for l in lines):
            reject('text beside a ruled table needs manual reconstruction.')
    for i, box in enumerate(boxes):
        if any(intersects(box, other) for other in boxes[i+1:]):
            reject('nested or overlapping table borders are not supported.')
        if any(intersects(box, r) or (box.y0 < r.y0 < box.y1 and r.x1 > box.x0 and r.x0 < box.x1) for r in rects+rules):
            reject('a table grid or graphics inside a bordered form need manual reconstruction.')
        if any(intersects(box, l.bbox) and not contains(box, l.bbox) and
               not blank_touches_border(box, l) for l in lines):
            reject('text crosses a form border.')
    # Join adjacent vector glyphs into one embedded logo. Text is never rasterized.
    clusters = []
    for rect in rects:
        if rect.is_empty:
            continue
        merged = rect
        again = True
        while again:
            again = False
            for old in list(clusters):
                if intersects(merged+(-10, -4, 10, 4), old):
                    merged |= old
                    clusters.remove(old)
                    again = True
        clusters.append(merged)
    for rect in clusters:
        if any(intersects(rect, l.bbox) for l in lines if l.text.strip()):
            reject('graphics overlap selectable text.')
        if rect.height > page.rect.height*.5:
            reject('a large graphic needs manual placement.')
    return boxes, rules, clusters, grids


def check_columns(lines):
    """Reject sustained column gaps, including fragments merged by PDF extraction."""
    blocks = {}
    for line in lines:
        blocks[line.block] = blocks.get(line.block, pymupdf.Rect(line.bbox)) | line.bbox
    values = list(blocks.values())
    for i, a in enumerate(values):
        for b in values[i+1:]:
            if min(a.y1, b.y1)-max(a.y0, b.y0) > 18 and (a.x1+8 < b.x0 or b.x1+8 < a.x0):
                reject('multiple body columns are not supported.')
    rows = {}
    for line in lines:
        if line.field or line.text.strip().startswith('('):
            continue
        baseline = round(line.spans[0]['origin'][1])
        rows.setdefault(baseline, []).extend((c, s['size']) for s in line.spans for c in s['chars'] if c['c'].strip())
    gaps = []
    for y, chars in rows.items():
        chars.sort(key=lambda pair: pair[0]['bbox'][0])
        for (a, size), (b, _) in zip(chars, chars[1:]):
            if b['bbox'][0]-a['bbox'][2] > max(24, size*3):
                gaps.append((y, a['bbox'][2], b['bbox'][0]))
    for y, x0, x1 in gaps:
        matching = {yy for yy, a, b in gaps if min(b, x1)-max(a, x0) > 12 and abs(yy-y) < 100}
        if len(matching) >= 3:
            reject('multiple body columns are not supported.')


def element(name, **attrs):
    node = OxmlElement('w:'+name)
    for key, value in attrs.items():
        node.set(qn('w:'+key), str(value))
    return node


def cell_margins(cell, value=0):
    pr = cell._tc.get_or_add_tcPr()
    for old in pr.findall(qn('w:tcMar')):
        pr.remove(old)
    margins = element('tcMar')
    for edge in ('top', 'left', 'bottom', 'right'):
        margins.append(element(edge, w=value, type='dxa'))
    pr.append(margins)


def cell_borders(cell, borders):
    pr = cell._tc.get_or_add_tcPr()
    for old in pr.findall(qn('w:tcBorders')):
        pr.remove(old)
    pr.insert_element_before(borders, 'w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection',
                             'w:tcFitText', 'w:vAlign', 'w:hideMark', 'w:headers',
                             'w:cellIns', 'w:cellDel', 'w:cellMerge', 'w:tcPrChange')


def table_at(parent, widths, border=False, border_style=(6, '000000')):
    if hasattr(parent, 'sections'):
        table = parent.add_table(rows=1, cols=len(widths))
    else:
        table = parent.add_table(rows=1, cols=len(widths))
        # Word requires a final paragraph in every cell. It should not add a
        # default-sized blank line after each nested form row.
        tail = parent.paragraphs[-1]
        tail.paragraph_format.space_after = Pt(0)
        tail.paragraph_format.line_spacing = Pt(.05)
        tail.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        tail.add_run().font.size = Pt(1)
    table.autofit = False
    pr = table._tbl.tblPr
    pr.append(element('tblInd', w=0, type='dxa'))
    borders = element('tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        borders.append(element(edge, val='single' if border and edge[:6] != 'inside' else 'nil',
                               sz=border_style[0], color=border_style[1]))
    pr.append(borders)
    for col, cell, width in zip(table.columns, table.rows[0].cells, widths):
        col.width = cell.width = Pt(width)
        cell_margins(cell, 50 if border else 0)
    return table


def split_continuation_rows(table):
    # Word otherwise moves a whole tall single-row form to the next page when
    # that row fits there. Paragraph rows keep a single flowing bordered table.
    first = table.cell(0, 0)
    children = [n for n in first._tc if n.tag != qn('w:tcPr')]
    for child in children:
        first._tc.remove(child)
    for i, child in enumerate(children):
        cell = first if i == 0 else table.add_row().cells[0]
        for old in list(cell._tc):
            if old.tag != qn('w:tcPr'):
                cell._tc.remove(old)
        cell_margins(cell, 0)
        margins = cell._tc.get_or_add_tcPr().find(qn('w:tcMar'))
        for edge in ('left', 'right'):
            margins.find(qn('w:'+edge)).set(qn('w:w'), '50')
        cell._tc.append(child)
        if child.tag != qn('w:p'):
            cell.add_paragraph()
            table_separator(cell)
    # Table-wide top/bottom borders repeat at every Word page boundary. Put
    # those edges on the actual first/last cells so continuation pages stay open.
    borders = table._tbl.tblPr.find(qn('w:tblBorders'))
    for edge, cell in [('top', table.cell(0, 0)), ('bottom', table.cell(len(table.rows)-1, 0))]:
        source = borders.find(qn('w:'+edge))
        own = cell._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
        if own is None:
            own = element('tcBorders')
            cell_borders(cell, own)
        own.append(element(edge, val='single', sz=source.get(qn('w:sz')), color=source.get(qn('w:color'))))
        source.set(qn('w:val'), 'nil')


def table_separator(parent):
    # Adjacent tables in Word merge into a single grid unless a paragraph
    # separates them; that changes the widths of unrelated form fields.
    p = parent.add_paragraph() if hasattr(parent, 'sections') else parent.paragraphs[-1]
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(.05)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.add_run().font.size = Pt(1)
    return p


class Writer:
    def __init__(self, doc, mapping, width, left, dpi, cancel, keep_line_breaks=False):
        self.doc, self.mapping = doc, mapping
        self.width, self.left, self.dpi, self.cancel = width, left, dpi, cancel
        self.expected = []
        self.field_count = 0
        self.last_field = None
        self.spacing = TextSpacing(mapping)
        self.keep_line_breaks = keep_line_breaks
        self.note_refs = {}
        self.note_rises = {}
        self.emitted_refs = set()
        self.emitted_squares = set()

    def minimum_spacing(self, lines):
        # Joining unequal-size source rows changes which font sets the first
        # Word baseline. The source-row descent model needs retained breaks.
        if not self.keep_line_breaks and max(l.size for l in lines)-min(l.size for l in lines) > .02:
            return None
        return self.spacing.minimum(lines)

    def text_metrics(self, para, lines, responsive=True):
        """Set paragraph-mark metrics too, including empty form/caption cells."""
        span = max((s for line in lines for s in line.spans), key=lambda s: s['size'])
        face = self.mapping[span['font']]
        pr = para._p.get_or_add_pPr()
        mark = element('rPr')
        mark.append(element('rFonts', ascii=face.family, hAnsi=face.family, eastAsia=face.family, cs=face.family))
        mark.append(element('b', val=int(face.bold)))
        mark.append(element('i', val=int(face.italic)))
        mark.append(element('sz', val=round(span['size']*2)))
        mark.append(element('szCs', val=round(span['size']*2)))
        pr.append(mark)
        multiple = self.spacing.multiple(lines) if responsive else None
        if multiple is not None:
            para.paragraph_format.line_spacing = multiple
        elif responsive:
            minimum = self.minimum_spacing(lines)
            if minimum is not None:
                para.paragraph_format.line_spacing = Pt(minimum[0])

    def run(self, para, span, text):
        face = self.mapping[span['font']]
        run = para.add_run('' if span.get('inline_square') else text)
        if span.get('inline_square'):
            if id(span) in self.emitted_squares:
                reject('an inline square was emitted more than once.')
            self.emitted_squares.add(id(span))
            data, width, height = span['inline_square']
            para.paragraph_format.line_spacing = Pt(max(height+span['size']*.3, span['size']*1.2))
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            run.add_picture(io.BytesIO(data), width=Pt(width), height=Pt(height))
        run.font.name = face.family
        fonts = run._r.get_or_add_rPr().rFonts
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:'+attr), face.family)
        run.bold, run.italic = face.bold, face.italic
        run.underline = bool(span.get('underline'))
        run.font.size = Pt(span['size'])
        run.font.color.rgb = RGBColor.from_string(f"{span['color']:06X}")
        if span.get('highlight_fill'):
            run._r.get_or_add_rPr().append(element('shd', val='clear', color='auto', fill=span['highlight_fill']))
        run._r.get_or_add_rPr().append(element('kern', val=1))
        if id(span) in self.note_rises:
            run._r.get_or_add_rPr().append(element('position', val=round(self.note_rises[id(span)]*2)))
        if id(span) in self.note_refs and text.strip() and id(span) not in self.emitted_refs:
            note = self.note_refs[id(span)]
            run._r.insert(1, element('footnoteReference', id=note.number, customMarkFollows=1))
            self.emitted_refs.add(id(span))
        self.expected.extend((c, face.family, face.bold, face.italic) for c in text if not c.isspace())
        return run

    def paragraph(self, parent, size, before=0):
        if not hasattr(parent, 'sections') and len(parent.paragraphs) == 1 and not parent.paragraphs[0].text and not len(parent.tables):
            para = parent.paragraphs[0]
        else:
            para = parent.add_paragraph()
        pf = para.paragraph_format
        pf.space_before = Pt(max(0, min(before, 18)))
        pf.space_after = Pt(0)
        pf.line_spacing = Pt(size*1.2)
        pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        pf.widow_control = True
        return para

    def text(self, parent, lines, before=0, title=False, left=None, width=None, measured_gap=False):
        self.last_field = None
        para = self.paragraph(parent, max(l.size for l in lines), before)
        first = lines[0]
        left = self.left if left is None else left
        width = self.width if width is None else width
        pf = para.paragraph_format
        if self.keep_line_breaks and measured_gap:
            # Only baseline/metric-derived gaps bypass the conservative cap.
            # A raw PDF bounding-box gap is not a Word paragraph distance.
            pf.space_before = Pt(max(0, before))
        starts = [l.bbox.x0 for l in lines]
        ends = [l.bbox.x1 for l in lines]
        centers = [(a+b)/2 for a, b in zip(starts, ends)]
        tolerance = max(2, first.size*.4)
        if title:
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif len(lines) == 1 and re.fullmatch(r'(?:[-–—.·=]\s*){5,}', first.text.strip()):
            # Decorative separators have no paragraph indent. Keeping their
            # literal characters in the full text width avoids a stray last dash.
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif (all(abs(c-(left+width/2)) < tolerance for c in centers) and
              all(l.bbox.width < width*.85 for l in lines) and
              (len(lines) > 1 or first.bold or first.text.lstrip().startswith('('))):
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif all(abs(x-(left+width)) < tolerance for x in ends) and min(starts)-left > width*.15:
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            # Reconstruct the continuation margin, then a first-line/hanging
            # indent. An indented address is not automatically a centered title.
            continuation = median(starts[1:]) if len(starts) > 1 else starts[0]
            inset = max(0, min(width*.8, continuation-left))
            pf.left_indent = Pt(inset)
            pf.first_line_indent = Pt(max(-inset, min(width*.4, starts[0]-continuation)))
            if len(lines) >= 3 and all(abs(x-(left+width)) < tolerance for x in ends[:-1]):
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if self.keep_line_breaks and len(lines) >= 2 and all(l.text.strip() for l in lines):
                text_ends = [max(c['bbox'][2] for s in l.spans for c in s['chars'] if not c['c'].isspace()) for l in lines]
                edge = median(text_ends[:-1])
                if (edge > left+width*.85 and all(abs(x-edge) < tolerance for x in text_ends[:-1])
                        and all(abs(x-continuation) < tolerance for x in starts[1:])):
                    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    # Round the source edge outward. Median/nearest rounding
                    # can make a full source line a fraction of a twip too wide.
                    pf.right_indent = Pt(max(0, math.floor((left+width-max(text_ends[:-1])-.1)*20)/20))
        if len(lines) > 1:
            baselines = [l.spans[0]['origin'][1] for l in lines]
            leading = median(b-a for a, b in zip(baselines, baselines[1:]))
            if first.size <= leading <= first.size*2:
                pf.line_spacing = Pt(leading)
        self.text_metrics(para, lines, responsive=not title)
        pf.keep_with_next = first.bold and len(lines) == 1
        for index, line in enumerate(lines):
            if index:
                # Keep hyphens, but never concatenate Cyrillic/Latin words merely
                # because PDF line boundaries are not actual space characters.
                previous = lines[index-1].text.rstrip()
                if self.keep_line_breaks:
                    self.run(para, line.spans[0], '').add_break()
                elif previous and not previous.endswith(('-', '\u00ad')):
                    self.run(para, line.spans[0], ' ')
            previous_x = None
            for si, span in enumerate(line.spans):
                text = ''.join(c['c'] for c in span['chars'])
                if si == 0:
                    text = text.lstrip()
                if si == len(line.spans)-1:
                    text = text.rstrip()
                if (not span.get('inline_square') and previous_x is not None and span['bbox'][0]-previous_x > span['size']*.65
                        and text and not text[0].isspace() and not previous_text.endswith(' ')):
                    self.run(para, span, ' ')
                tracking = (source_tracking(span, self.mapping[span['font']])
                            if self.keep_line_breaks and id(span) not in self.note_rises and not span.get('inline_square') else None)
                if tracking is not None:
                    raw = ''.join(c['c'] for c in span['chars'])
                    start = len(raw)-len(raw.lstrip()) if si == 0 else 0
                    end = len(raw.rstrip()) if si == len(line.spans)-1 else len(raw)
                    groups = []
                    for c, pitch in zip(raw[start:end], tracking[start:end]):
                        c = ' ' if c.isspace() else c
                        if c == ' ' and groups and groups[-1][0].endswith(' '):
                            continue
                        if groups and groups[-1][1] == pitch:
                            groups[-1] = (groups[-1][0]+c, pitch)
                        else:
                            groups.append((c, pitch))
                    for part, pitch in groups:
                        run = self.run(para, span, part)
                        pr = run._r.get_or_add_rPr()
                        pr.remove(pr.find(qn('w:kern')))
                        pr.append(element('spacing', val=pitch))
                else:
                    self.run(para, span, re.sub(r'\s+', ' ', text))
                previous_text = text
                previous_x = span['bbox'][2]
        return para

    def field(self, parent, line, width):
        # Every label/value portion gets an expanding cell. Replacing a blank
        # with several lines grows the row and moves following content naturally.
        chars = [(c, s) for s in line.spans for c in s['chars']]
        text = ''.join(c['c'] for c, _ in chars)
        matches = blank_ranges(text)
        parts, pos = [], 0
        for start, end in matches:
            if start > pos:
                parts.append((chars[pos:start], False))
            parts.append((chars[start:end], True))
            pos = end
        if pos < len(chars):
            parts.append((chars[pos:], False))
        # Keep source boundaries, including real whitespace between fields.
        # Padding every label by a fraction of its font size moves later fields
        # and consumes the blanks. Only a twip-rounding allowance is needed.
        edges = [cs[0][0]['bbox'][0] for cs, _ in parts]+[chars[-1][0]['bbox'][2]]
        for i in range(1, len(parts)):
            if parts[i-1][1]:
                edges[i] = parts[i-1][0][-1][0]['bbox'][2]
        origin = edges[0]
        edges = [x-edges[0] for x in edges]
        for i in range(1, len(parts)):
            if not parts[i-1][1] and parts[i][1]:
                edges[i] += .15
        edges[-1] += .15 if not parts[-1][1] else 0
        scale = min(1, width/edges[-1])
        rounded = [round(x*scale*20) for x in edges]
        widths = [(b-a)/20 for a, b in zip(rounded, rounded[1:])]
        if any(v < .5 for v in widths):
            reject('a form row has insufficient space for editable values.')
        table = table_at(parent, widths)
        for i, (cell, (cs, blank)) in enumerate(zip(table.rows[0].cells, parts)):
            para = self.paragraph(cell, line.size)
            self.text_metrics(para, [line])
            if blank:
                self.field_count += 1
                borders = element('tcBorders')
                borders.append(element('bottom', val='single', sz=4, color='000000'))
                cell_borders(cell, borders)
                self.run(para, cs[0][1], ' ')
            else:
                # A real PDF gap after a blank belongs before the next label,
                # not under the blank's rule. Keep it as a paragraph inset.
                gap = (cs[0][0]['bbox'][0]-origin)*scale-rounded[i]/20
                if gap > .05:
                    para.paragraph_format.left_indent = Pt(gap)
                current, text = cs[0][1], ''
                for char, span in cs:
                    if span is not current:
                        self.run(para, current, text)
                        current, text = span, ''
                    text += char['c']
                self.run(para, current, text)
        separator = table_separator(parent)
        self.last_field = (widths, parts, line, table, separator)
        return table

    def caption(self, parent, line):
        if not self.last_field:
            return False
        chars = [(c, s) for s in line.spans for c in s['chars']]
        text = ''.join(c['c'] for c, _ in chars)
        matches = list(re.finditer(r'\([^()]*\)', text))
        if not matches or re.sub(r'\([^()]*\)', '', text).strip():
            return False
        widths, parts, field_line, field_table, separator = self.last_field
        if not -line.size < line.bbox.y0-field_line.bbox.y1 < line.size:
            return False
        # The value and its caption are one field, even at a page boundary.
        for cell in field_table.rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
        separator.paragraph_format.keep_with_next = True
        blanks = [(i, (cs[0][0]['bbox'][0]+cs[-1][0]['bbox'][2])/2)
                  for i, (cs, blank) in enumerate(parts) if blank]
        table = table_at(parent, widths)
        used = set()
        for match in matches:
            cs = chars[match.start():match.end()]
            center = (cs[0][0]['bbox'][0]+cs[-1][0]['bbox'][2])/2
            choices = [(abs(center-x), i) for i, x in blanks if i not in used]
            if not choices:
                reject('form captions cannot be assigned to distinct fields.')
            _, index = min(choices)
            used.add(index)
            p = self.paragraph(table.cell(0, index), line.size)
            self.text_metrics(p, [line])
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            span, text = cs[0][1], ''
            for char, current in cs:
                if current is not span:
                    self.run(p, span, text)
                    span, text = current, ''
                text += char['c']
            self.run(p, span, text)
        # Empty cells must have the same small, expanding paragraph metrics.
        for i, cell in enumerate(table.rows[0].cells):
            if i not in used:
                self.text_metrics(self.paragraph(cell, line.size), [line])
        table_separator(parent)
        self.last_field = None
        return True

    def image(self, parent, page, rect, width=None):
        para = self.paragraph(parent, 1)
        pix = page.get_pixmap(matrix=pymupdf.Matrix(self.dpi/72, self.dpi/72), clip=rect+(-.5, -.5, .5, .5), alpha=False)
        para.add_run().add_picture(io.BytesIO(pix.tobytes('png')), width=Pt(min(width or rect.width, self.width)))

    def grid(self, parent, grid, cells):
        inset = max(0, grid.xs[0]-self.left)
        scale = min(1, (self.width-inset)/grid.bbox.width)
        widths = [(b-a)*scale for a, b in zip(grid.xs, grid.xs[1:])]
        table = table_at(parent, widths)
        table._tbl.tblPr.find(qn('w:tblInd')).set(qn('w:w'), str(round(inset*20)))
        for _ in cells[1:]:
            table.add_row()
        for r, row in enumerate(cells):
            for c, lines in enumerate(row):
                cell = table.cell(r, c)
                cell.width = Pt(widths[c])
                cell_margins(cell, 50)
                borders = element('tcBorders')
                for edge, (size, color) in [('top', grid.horizontal[r]), ('bottom', grid.horizontal[r+1]),
                                             ('left', grid.vertical[c]), ('right', grid.vertical[c+1])]:
                    borders.append(element(edge, val='single', sz=size, color=color))
                cell_borders(cell, borders)
                if lines:
                    self.flow(cell, lines, widths[c]-5, grid.ys[r], left=grid.xs[c]+2.5)
                else:
                    self.paragraph(cell, 8)
        table_separator(parent)
        return table

    def flow(self, parent, lines, width, top=None, left=None):
        lines = list(lines)
        # Extractors sometimes return a whitespace-only row below the last
        # signature. It must not create an otherwise empty final Word page.
        while lines and not lines[-1].text.strip():
            lines.pop()
        check_columns(lines)
        self.spacing.flow(lines)
        left = self.left if left is None else left
        self.last_field = None
        pending = []
        previous = None
        before = 0
        previous_group = None
        emitted = None
        def flush():
            nonlocal pending, emitted
            if pending:
                spacing = before
                measured_gap = False
                leading = self.spacing.leading(pending)
                minimum = self.minimum_spacing(pending) if self.spacing.multiple(pending) is None else None
                if (previous_group is not None and leading is not None and emitted is not None
                        and emitted[0] is previous_group and emitted[1] is not None
                        and self.spacing.style([previous_group, pending[0]]) == self.spacing.style(pending)):
                    delta = pending[0].spans[0]['origin'][1]-previous_group.spans[0]['origin'][1]
                    spacing = max(0, delta-emitted[1])
                    measured_gap = True
                elif (previous_group is not None and emitted is not None and emitted[0] is previous_group
                      and minimum is not None and emitted[2] is not None and minimum[3] == emitted[2][3]):
                    delta = pending[0].spans[0]['origin'][1]-previous_group.spans[0]['origin'][1]
                    spacing = max(0, delta-minimum[0]-emitted[2][2]+minimum[1])
                    measured_gap = True
                self.text(parent, pending, spacing, left=left, width=width, measured_gap=measured_gap)
                emitted = (pending[-1], leading, minimum)
                pending = []
        for line in lines:
            cancelled(self.cancel)
            gap = line.bbox.y0-(previous.bbox.y1 if previous else (top if top is not None else line.bbox.y0))
            boundary = (previous is None or line.block != previous.block or line.field or previous.field or
                        id(line) in self.spacing.paragraph_starts or
                        (previous is not None and abs(line.size-previous.size) > min(line.size, previous.size)*.1
                         and all(abs(s['size']-line.size) < .02 for s in line.spans)
                         and all(abs(s['size']-previous.size) < .02 for s in previous.spans)) or
                        (previous.bold and not line.bold) or
                        bool(re.match(r'^\s*(?:[-•·]|\d+[.)])\s', line.text)) or
                        (previous is not None and bool(re.match(r'^\s*[-•·]\s', previous.text)) and
                         previous.bbox.x1 < left+width-line.size*4) or
                        gap > line.size*.6 or
                        (previous.bbox.x1 < left+width-max(line.size*4, width*.12) and
                         previous.bbox.x0 < left+width*.1 and
                         (previous.text.rstrip().endswith(('.', ':', ';', ')')) or
                          line.text.lstrip()[:1].isupper() or re.match(r'^\s*[\w -]{2,35}:', line.text))))
            if boundary:
                flush()
                previous_group = previous if previous is not None and not previous.field else None
                before = max(0, gap)
            if line.field:
                if (previous is not None and not previous.field and
                        previous.bbox.width < width*.6 and
                        not previous.text.lstrip().startswith('(') and parent.paragraphs):
                    # Short labels immediately above a signature/value belong
                    # with that field rather than at the foot of the prior page.
                    parent.paragraphs[-1].paragraph_format.keep_with_next = True
                if before > 1 and parent.paragraphs:
                    parent.paragraphs[-1].paragraph_format.space_after = Pt(min(before, 12))
                self.field(parent, line, width)
            elif self.caption(parent, line):
                pass
            else:
                # A later parenthetical is narrative, not the earlier field's
                # caption. Do this when queuing text, before deferred emission.
                self.last_field = None
                pending.append(line)
            previous = line
        flush()


def build_document(pdf, mapping, dpi=144, cancel=None, keep_line_breaks=False):
    with prepared_pdf(pdf, reject) as (clean, highlights):
        return _build_document(clean, mapping, dpi, cancel, keep_line_breaks, highlights)


def _build_document(pdf, mapping, dpi, cancel, keep_line_breaks, highlights):
    pages = []
    continuations = continuation_borders(pdf)
    continuation_rects = {(n, tuple(pdf[n].get_drawings()[i]["rect"])): value for (n,i),value in continuations.items()}
    border_styles = {(p.number, tuple(d['rect'])):
                     (max(2, min(96, round(d['width']*8))), ''.join(f'{round(c*255):02X}' for c in d['color']))
                     for p in pdf for d in p.get_drawings()
                     if d['type'] == 's' and d.get('color') is not None and len(d['color']) == 3}
    note_plans, notes = [], []
    for page in pdf:
        cancelled(cancel)
        lines = lines_on(page)
        if highlights.get(page.number):
            apply_highlights(lines, highlights[page.number], reject)
        boxes, rules, graphics, grids = graphics_on(page, lines, continuations)
        plan = find_notes(lines, rules, page.rect.height, reject)
        note_plans.append(plan)
        for note in plan[0]:
            note.number = len(notes)+1
            notes.append(note)
        tables = [(grid, *partition(grid, lines, reject)) for grid in grids]
        consumed = {item for _, _, ids in tables for item in ids}
        sequence = [(l.bbox.y0, [l]) for l in lines if id(l) not in consumed]
        sequence += [(g.ys[0], [l for row in cells for cell in row for l in cell]) for g, cells, _ in tables]
        reading = [l for _, group in sorted(sequence, key=lambda pair: pair[0]) for l in group]
        pages.append((lines, boxes, rules, graphics, tables, reading))
    all_lines = [l for lines, *_ in pages for l in lines if l.text.strip()]
    left = min(72, max(12, min(l.bbox.x0 for l in all_lines)))
    right = min(72, max(12, pdf[0].rect.width-max(l.bbox.x1 for l in all_lines)))
    if keep_line_breaks:
        # A full source row can touch the inferred page margin. Round the
        # available area outward, including small source margins, rather than
        # forcing its last word onto an extra line because of twip rounding.
        left = math.floor(min(72, max(0, min(l.bbox.x0 for l in all_lines)))*20)/20
        right = math.floor(min(72, max(0, pdf[0].rect.width-max(l.bbox.x1 for l in all_lines)-.1))*20)/20
    width = pdf[0].rect.width-left-right
    doc = Document()
    default_span = max((s for l in all_lines for s in l.spans), key=lambda s: len(s['chars']))
    normal = doc.styles['Normal']
    normal.font.name = mapping[default_span['font']].family
    normal.font.size = Pt(default_span['size'])
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(pdf[0].rect.width), Pt(pdf[0].rect.height)
    sec.left_margin, sec.right_margin = Pt(left), Pt(right)
    sec.top_margin = sec.bottom_margin = Pt(18)
    writer = Writer(doc, mapping, width, left, dpi, cancel, keep_line_breaks)
    writer.note_refs = {id(n.anchor): n for n in notes}
    writer.note_rises = {id(s): rise for n in notes for s, rise in [(n.anchor, n.rise), (n.label, n.label_rise)]}
    if keep_line_breaks:
        for node in list(doc.settings.element.iter(qn('w:doNotExpandShiftReturn'))):
            node.getparent().remove(node)
    source_chars = Counter(c['c'] for lines, *_ in pages for l in lines for s in l.spans for c in s['chars'] if not c['c'].isspace())
    removed = Counter()
    continuing_tables = {}
    for page, (lines, boxes, rules, graphics, tables, reading) in zip(pdf, pages):
        # Continue ordinary body pages. A footer note is tied to its source-page
        # anchor: retain the following boundary so the next page's body does not
        # get pulled above it. This is a normal editable Word page break.
        body_size = sorted(l.size for l in lines)[len(lines)//2]
        new_title = lines[0].bold and lines[0].size > body_size*1.2
        if page.number and (note_plans[page.number-1][0] or new_title or
                            any(g.y0 < page.rect.height*.15 for g in graphics)):
            doc.add_page_break()
        _, note_ids, note_rules = note_plans[page.number]
        consumed = set(note_ids)
        events = []
        for grid, cells, ids in tables:
            consumed.update(ids)
            events.append((grid.ys[0], 'grid', (grid, cells)))
        for graphic in graphics:
            title = [l for l in lines if l.bbox.x0 > graphic.x1+5 and
                     l.bbox.y0 < graphic.y1+20 and l.bbox.y1 > graphic.y0]
            side = 'left'
            if not title and graphic.x0 > page.rect.width/2:
                adjacent = [l for l in lines if l.bbox.x1 < graphic.x0-5 and
                            l.bbox.y0 < graphic.y1 and l.bbox.y1 > graphic.y0]
                # A heading and its subordinate date/subtitle can share a
                # source block beside a right-hand logo. Do not absorb later
                # body blocks or turn arbitrary side-by-side text into a header.
                if (adjacent and adjacent[0].bold and len({l.block for l in adjacent}) == 1
                        and all(l.size <= adjacent[0].size for l in adjacent)):
                    title, side = adjacent, 'right'
            header = (graphic.y0 < page.rect.height*.15 and graphic.height < page.rect.height*.15 and
                      title and (side == 'right' or all(l.bold for l in title)))
            if header:
                consumed.update(id(l) for l in title)
                events.append((min(graphic.y0, title[0].bbox.y0), 'header', (graphic, title, side)))
            else:
                events.append((graphic.y0, 'image', graphic))
        for box in boxes:
            inner = [l for l in lines if contains(box, l.bbox)]
            if not inner:
                reject('an empty drawn box is not a supported editable field.')
            consumed.update(id(l) for l in inner)
            events.append((box.y0, 'box', (box, inner)))
        body = [l for l in lines if id(l) not in consumed]
        check_columns(body)
        for line in body:
            events.append((line.bbox.y0, 'line', line))
        for rule in rules:
            if id(rule) in note_rules:
                continue
            underline = [s for l in lines for s in l.spans
                         if abs(s['origin'][1]-rule.y0) < s['size']*.35 and
                         abs(s['bbox'][0]-rule.x0) < 2 and abs(s['bbox'][2]-rule.x1) < 2]
            if underline:
                for span in underline:
                    span['underline'] = True
                continue
            if any(intersects(l.bbox, rule+(0, -.5, 0, .5), .1) for l in lines):
                reject('a drawn line crosses text.')
            events.append((rule.y0, 'rule', rule))
        pending, last_bottom = [], None
        def flush_body():
            nonlocal pending, last_bottom
            if pending:
                writer.flow(doc, pending, width, last_bottom)
                last_bottom = pending[-1].bbox.y1
                pending = []
        for _, kind, data in sorted(events, key=lambda e: e[0]):
            if kind == 'line':
                pending.append(data)
                continue
            flush_body()
            if kind == 'header':
                graphic, title, side = data
                if side == 'right':
                    split = graphic.x0-left
                    # A logo can extend beyond the body text's right edge.
                    # Preserve its actual width in this header only.
                    table = table_at(doc, [split, graphic.width])
                    writer.flow(table.cell(0, 0), title, split, min(graphic.y0, title[0].bbox.y0), left=left)
                    writer.image(table.cell(0, 1), page, graphic, graphic.width)
                else:
                    split = min(width*.45, max(graphic.width+8, min(l.bbox.x0 for l in title)-left-6))
                    table = table_at(doc, [split, width-split])
                    writer.image(table.cell(0, 0), page, graphic, min(graphic.width, split-4))
                    # Each source title line remains a paragraph, not a positioned box.
                    for line in title:
                        writer.text(table.cell(0, 1), [line], title=True)
                table_separator(doc)
                last_bottom = max(graphic.y1, title[-1].bbox.y1)
            elif kind == 'box':
                box, inner = data
                if doc.paragraphs:
                    doc.paragraphs[-1].paragraph_format.space_after = Pt(3)
                continuation = continuation_rects.get((page.number, tuple(box)))
                if continuation and continuation[0] == 'end':
                    table = continuing_tables.pop(continuation[1])
                else:
                    table = table_at(doc, [box.width if continuation else width], border=True,
                                     border_style=border_styles.get((page.number, tuple(box)), (6, '000000')))
                    if continuation:
                        table._tbl.tblPr.find(qn('w:tblInd')).set(qn('w:w'), str(round((box.x0-left)*20)))
                    if continuation:
                        continuing_tables[continuation[1]] = table
                writer.flow(table.cell(0, 0), inner, (box.width if continuation else width)-5, box.y0, left=box.x0+2.5)
                if continuation and continuation[0] == 'end':
                    split_continuation_rows(table)
                table_separator(doc)
                last_bottom = box.y1
            elif kind == 'image':
                writer.image(doc, page, data)
                last_bottom = data.y1
            elif kind == 'grid':
                grid, cells = data
                writer.grid(doc, grid, cells)
                last_bottom = grid.ys[-1]
            else:
                table = table_at(doc, [min(width, data.width)])
                borders = element('tcBorders')
                borders.append(element('bottom', val='single', sz=4, color='000000'))
                cell_borders(table.cell(0, 0), borders)
                writer.paragraph(table.cell(0, 0), 7.5)
                table_separator(doc)
                last_bottom = data.y1
        flush_body()
        for line in lines:
            for start, end in blank_ranges(line.text):
                removed.update(line.text[start:end])
    if continuing_tables:
        reject('a continued form has no matching end.')
    expected_squares = {id(s) for lines, *_ in pages for l in lines for s in l.spans if s.get('inline_square')}
    if writer.emitted_squares != expected_squares:
        reject('inline square conservation failed.')
    body_expected = writer.expected
    note_expected = {}
    if notes:
        part = create_part(doc)
        for note in notes:
            writer.expected = []
            writer.flow(add_note(part, note.number), note.lines, width, left=left)
            note_expected[str(note.number)] = writer.expected
        writer.expected = body_expected
    actual_chars = Counter(c for c, *_ in body_expected)
    actual_chars.update(c for signature in note_expected.values() for c, *_ in signature)
    if source_chars != actual_chars+removed:
        raise ValueError('Editable conversion did not account for every source character.')
    # Independently compare the emitted order and face of every non-whitespace
    # character with the extracted reading sequence, rather than trusting only
    # a ledger written by the emitter itself.
    def signature(lines, anchors=None):
        result = []
        for line in lines:
            blanks = {i for start, end in blank_ranges(line.text) for i in range(start, end)}
            index = 0
            for span in line.spans:
                face = mapping[span['font']]
                if anchors is not None and id(span) in writer.note_refs:
                    note = writer.note_refs[id(span)]
                    anchors.append((len(result), note.marker, str(note.number)))
                for char in span['chars']:
                    if index not in blanks and not char['c'].isspace():
                        result.append((char['c'], face.family, face.bold, face.italic))
                    index += 1
        return result
    anchors = []
    reading = [line for page, plan in zip(pages, note_plans) for line in page[-1] if id(line) not in plan[1]]
    source_signature = signature(reading, anchors)
    source_notes = {str(n.number): signature(n.lines) for n in notes}
    if body_expected != source_signature or note_expected != source_notes:
        raise ValueError('Editable output text, reading order or emphasis verification failed.')
    expected = {'body': source_signature, 'notes': source_notes, 'anchors': anchors} if notes else source_signature
    return doc, expected, writer.field_count


def verify_document(doc, expected):
    def check_part(root):
        if list(root.iter(qn('w:txbxContent'))) or list(root.iter(qn('w:framePr'))):
            raise ValueError('Editable output unexpectedly contains positioned text.')
        actual, anchors = [], []
        for run in root.iter(qn('w:r')):
            texts = ''.join(t.text or '' for t in run.iter(qn('w:t')))
            for ref in run.findall(qn('w:footnoteReference')):
                if ref.get(qn('w:customMarkFollows')) != '1' or not re.fullmatch(r'(?:[0-9]{1,3}|[*†‡])', texts):
                    raise ValueError('Editable footnote reference verification failed.')
                anchors.append((len(actual), texts, ref.get(qn('w:id'))))
            if not texts.strip():
                continue
            rpr = run.find(qn('w:rPr'))
            fonts = rpr.find(qn('w:rFonts')) if rpr is not None else None
            if fonts is None:
                raise ValueError('Editable output contains text without an approved font.')
            def on(name):
                node = rpr.find(qn('w:'+name))
                return node is not None and node.get(qn('w:val'), '1') not in ('0', 'false')
            actual.extend((c, fonts.get(qn('w:ascii')), on('b'), on('i')) for c in texts if not c.isspace())
        return actual, anchors
    actual, anchors = check_part(doc.element)
    if isinstance(expected, dict):
        root = read_part(doc)
        actual_notes = {}
        for node in root.findall(qn('w:footnote')):
            number = node.get(qn('w:id'))
            if node.get(qn('w:type')) in ('separator', 'continuationSeparator') and number in ('-1', '0'):
                if list(node.iter(qn('w:t'))):
                    raise ValueError('Editable footnote separator verification failed.')
                continue
            if number in actual_notes or number not in expected['notes']:
                raise ValueError('Editable footnote identity verification failed.')
            actual_notes[number], nested = check_part(node)
            if nested:
                raise ValueError('Editable footnote reference verification failed.')
        if actual_notes != expected['notes'] or anchors != expected['anchors']:
            raise ValueError('Editable footnote content or reference verification failed.')
        expected = expected['body']
    elif anchors or any(r.reltype.endswith('/footnotes') for r in doc.part.rels.values()):
        raise ValueError('Editable output contains unexpected footnotes.')
    if actual != expected:
        raise ValueError('Editable output text, reading order or emphasis verification failed.')
